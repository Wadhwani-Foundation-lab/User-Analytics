"""
Runner: reads NEP_Test_Report.xlsx, adds 7 LLM-judge columns, saves NEP_Judge_Report.xlsx.

The judge evaluates the SQL and response for every row in the "All Questions" sheet.
New columns appended after the last existing column:
  N  Judge Verdict        — ACCURATE | INACCURATE | UNCERTAIN | SKIPPED
  O  Judge Confidence %   — 0-100
  P  Reasoning            — detailed explanation
  Q  SQL Issues           — specific problems with the SQL (empty if fine)
  R  Data Accuracy Note   — whether numbers in the response are plausible
  S  Recommended SQL      — corrected / optimized SQL (empty if original is fine)
  T  Recommended Response — better response wording (empty if original is fine)

SQL and Answer columns are written in full by the test suite runner (no truncation).
"""
from __future__ import annotations

import time
from pathlib import Path
from typing import Optional

from .config import DOCS_DIR

DEFAULT_INPUT = DOCS_DIR / "NEP_Test_Report.xlsx"
DEFAULT_OUTPUT = DOCS_DIR / "NEP_Judge_Report.xlsx"

JUDGE_COLS = [
    "Judge Verdict",
    "Judge Confidence %",
    "Reasoning",
    "SQL Issues",
    "Data Accuracy Note",
    "Recommended SQL",
    "Recommended Response",
]

JUDGE_COL_WIDTHS = [18, 18, 70, 50, 55, 80, 80]

# Colour palette — matches test suite report palette for consistency
C_HEADER_BG   = "1E3A5F"
C_HEADER_FG   = "FFFFFF"
C_ACCURATE    = "D4EDDA"   # green
C_INACCURATE  = "F8D7DA"   # red
C_UNCERTAIN   = "FFF3CD"   # yellow
C_SKIPPED     = "EBEBEB"   # grey


def run(
    input_path: Optional[Path] = None,
    output_path: Optional[Path] = None,
    limit: Optional[int] = None,
    delay_secs: float = 1.0,
    verbose: bool = False,
) -> None:
    import openpyxl
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    from .judge import evaluate

    input_path = input_path or DEFAULT_INPUT
    output_path = output_path or DEFAULT_OUTPUT

    if not input_path.exists():
        print(f"  ERROR: Input file not found: {input_path}")
        return

    print(f"\n{'='*64}")
    print(f"  NEP Analytics — LLM-as-Judge  (Claude Opus 4.8)")
    print(f"  Input:  {input_path.name}")
    print(f"  Output: {output_path.name}")
    print(f"{'='*64}\n")

    wb = openpyxl.load_workbook(input_path)

    if "All Questions" not in wb.sheetnames:
        print("  ERROR: 'All Questions' sheet not found in the input file.")
        return

    ws = wb["All Questions"]

    # Row 1 is the title merge, row 2 is headers, data starts at row 3.
    HEADER_ROW = 2

    # Build header → column index map
    headers = {
        ws.cell(row=HEADER_ROW, column=c).value: c
        for c in range(1, ws.max_column + 1)
    }

    col_num      = headers.get("#")
    col_question = headers.get("Question")
    col_path     = headers.get("Path")
    col_metric   = headers.get("Certified Metric")
    col_answer   = headers.get("Answer")
    col_sql      = headers.get("SQL Used")
    col_rows     = headers.get("Rows")

    # Append judge columns after last existing column
    start_col = ws.max_column + 1

    def _hdr_fill(hex_color: str) -> PatternFill:
        return PatternFill("solid", fgColor=hex_color)

    def _cell_fill(hex_color: str) -> PatternFill:
        return PatternFill("solid", fgColor=hex_color)

    # Write judge column headers
    for i, col_name in enumerate(JUDGE_COLS):
        c = ws.cell(row=HEADER_ROW, column=start_col + i, value=col_name)
        c.font = Font(name="Calibri", bold=True, color=C_HEADER_FG, size=10)
        c.fill = _hdr_fill(C_HEADER_BG)
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        ws.column_dimensions[get_column_letter(start_col + i)].width = JUDGE_COL_WIDTHS[i]

    # Determine how many data rows to process
    total_rows = ws.max_row - HEADER_ROW  # subtract title + header
    if limit:
        total_rows = min(total_rows, limit)

    print(f"  Evaluating {total_rows} question(s)...\n")

    accurate = inaccurate = uncertain = skipped = errors = 0
    verdict_color_map = {
        "ACCURATE":   C_ACCURATE,
        "INACCURATE": C_INACCURATE,
        "UNCERTAIN":  C_UNCERTAIN,
        "SKIPPED":    C_SKIPPED,
    }

    for offset in range(total_rows):
        row = offset + 3  # data rows start at row 3

        q_num    = ws.cell(row=row, column=col_num).value      if col_num      else offset + 1
        question = ws.cell(row=row, column=col_question).value if col_question else ""
        path     = ws.cell(row=row, column=col_path).value     if col_path     else ""
        metric   = ws.cell(row=row, column=col_metric).value   if col_metric   else None
        answer   = ws.cell(row=row, column=col_answer).value   if col_answer   else ""
        sql_raw  = ws.cell(row=row, column=col_sql).value      if col_sql      else ""
        rc_raw   = ws.cell(row=row, column=col_rows).value     if col_rows     else None

        question = str(question or "")
        sql_used = str(sql_raw or "")
        answer   = str(answer  or "")
        path     = str(path    or "")
        row_count = int(rc_raw) if rc_raw is not None and str(rc_raw).strip().isdigit() else None

        label = question[:68] + ("…" if len(question) > 68 else "")
        print(f"  [{offset+1:>3}/{total_rows}] Q{q_num}: {label}", end="", flush=True)

        t0 = time.monotonic()
        result = evaluate(
            question=question,
            sql_used=sql_used,
            answer=answer,
            row_count=row_count,
            path=path,
            metric=str(metric) if metric else None,
        )
        elapsed = time.monotonic() - t0

        if result.verdict == "ACCURATE":   accurate  += 1
        elif result.verdict == "INACCURATE": inaccurate += 1
        elif result.verdict == "SKIPPED":  skipped   += 1
        else:                               uncertain += 1
        if result.error:                    errors    += 1

        vcolor = verdict_color_map.get(result.verdict, C_UNCERTAIN)
        values = [
            result.verdict,
            result.confidence_pct,
            result.reasoning,
            result.sql_issues,
            result.data_accuracy_note,
            result.recommended_sql,
            result.recommended_response,
        ]

        for i, val in enumerate(values):
            c = ws.cell(row=row, column=start_col + i, value=val or "")
            c.font = Font(name="Calibri", bold=(i == 0), size=10)
            c.alignment = Alignment(horizontal="left", vertical="top", wrap_text=True)
            if i == 0:
                c.fill = _cell_fill(vcolor)

        if verbose and result.sql_issues:
            print(f"\n       SQL issues: {result.sql_issues[:120]}")
        if verbose and result.reasoning:
            print(f"       Reasoning:  {result.reasoning[:120]}")

        print(f"  → {result.verdict} ({result.confidence_pct}%)  {elapsed:.1f}s")

        if delay_secs > 0 and offset < total_rows - 1:
            time.sleep(delay_secs)

    wb.save(output_path)

    total_eval = accurate + inaccurate + uncertain + skipped
    print(f"\n{'='*64}")
    print(f"  ACCURATE:    {accurate:>4}  ({round(accurate/total_eval*100,1) if total_eval else 0}%)")
    print(f"  INACCURATE:  {inaccurate:>4}  ({round(inaccurate/total_eval*100,1) if total_eval else 0}%)")
    print(f"  UNCERTAIN:   {uncertain:>4}  ({round(uncertain/total_eval*100,1) if total_eval else 0}%)")
    print(f"  SKIPPED:     {skipped:>4}  (no SQL)")
    if errors:
        print(f"  API ERRORS:  {errors:>4}")
    print(f"\n  Judge report → {output_path}")
    print(f"{'='*64}\n")
