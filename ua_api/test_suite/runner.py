"""
Test suite runner for NEP Analytics.

Reads questions from docs/NEP Test Questions.xlsx, fires each against the
chat_adapter API, and writes a full report to docs/NEP_Test_Report.xlsx.

Also runs the multi-turn conversation scenarios from multiturn_scenarios.py.

Usage:
  python -m ua_api.test_suite                    # all questions
  python -m ua_api.test_suite --limit 10         # first 10 questions (smoke test)
  python -m ua_api.test_suite --category "Multi-Filter Mentor Network Queries"
  python -m ua_api.test_suite --skip-multiturn   # single-turn questions only
  python -m ua_api.test_suite --output my_report.xlsx
"""
from __future__ import annotations

import os
import sys
import time
from dataclasses import dataclass, field
from datetime import datetime
from pathlib import Path
from typing import Any, Dict, List, Optional

import requests
from dotenv import load_dotenv

# ── Path setup ────────────────────────────────────────────────────────────────

REPO_ROOT = Path(__file__).resolve().parent.parent.parent
DOCS_DIR = REPO_ROOT / "docs"
INPUT_XLSX = DOCS_DIR / "NEP Test Questions.xlsx"
DEFAULT_OUTPUT = DOCS_DIR / "NEP_Test_Report.xlsx"

# Load API key from chat_api/.env
_DEFAULT_ENV = REPO_ROOT / "chat_api" / ".env"
if _DEFAULT_ENV.exists():
    load_dotenv(_DEFAULT_ENV, override=False)
load_dotenv(override=False)

BASE_URL = os.getenv("ADAPTER_URL", "http://localhost:8001")
API_KEY = os.getenv("API_SECRET_KEY", "")
TIMEOUT = int(os.getenv("TEST_TIMEOUT", "90"))

# ── Result types ──────────────────────────────────────────────────────────────

STATUS_PASS = "PASS"
STATUS_ERROR = "ERROR"
STATUS_PARTIAL = "PARTIAL"   # responded but answer contains a DB/SQL error


@dataclass
class SingleTurnResult:
    number: int
    category: str
    question: str
    status: str
    path: str                    # "certified" | "text_to_sql" | "error"
    metric: Optional[str]
    confidence: Optional[float]
    response_type: Optional[str]
    answer: str
    sql_used: str
    row_count: Optional[int]
    latency_ms: int
    error: str


@dataclass
class MultiTurnResult:
    scenario_id: int
    scenario_name: str
    scenario_description: str
    turn_number: int
    question: str
    expected_intent: str
    expected_path: Optional[str]
    status: str
    path: str
    metric: Optional[str]
    confidence: Optional[float]
    response_type: Optional[str]
    answer: str
    sql_used: str
    latency_ms: int
    error: str
    context_turns: int           # how many prior turns were in the session


# ── API call ──────────────────────────────────────────────────────────────────

def _call_api(
    session_id: str,
    question: str,
    history: List[Dict[str, str]],
) -> Dict[str, Any]:
    headers = {"Content-Type": "application/json"}
    if API_KEY:
        headers["x-api-key"] = API_KEY
    payload = {"session_id": session_id, "question": question, "history": history}
    resp = requests.post(
        f"{BASE_URL}/api/chat",
        json=payload,
        headers=headers,
        timeout=TIMEOUT,
    )
    resp.raise_for_status()
    return resp.json()


def _classify_status(data: Dict[str, Any]) -> str:
    answer = (data.get("answer") or "").lower()
    if any(kw in answer for kw in [
        "database query failed", "sql error", "error executing",
        "could not execute", "column does not exist", "syntax error",
    ]):
        return STATUS_PARTIAL
    return STATUS_PASS


def _extract_result(data: Dict[str, Any]) -> Dict[str, Any]:
    prov = data.get("provenance") or {}
    table = data.get("table_data") or {}
    rows = table.get("rows") if isinstance(table, dict) else None
    row_count = len(rows) if rows is not None else None
    return {
        "path": "certified" if prov.get("certified") else "text_to_sql",
        "metric": prov.get("metric"),
        "confidence": prov.get("confidence"),
        "response_type": data.get("response_type"),
        "answer": data.get("answer", ""),
        "sql_used": data.get("sql_used") or "",
        "row_count": row_count,
    }


# ── Single-turn runner ────────────────────────────────────────────────────────

def run_single_turn(
    questions: List[Dict[str, Any]],
    session_prefix: str = "test-single",
) -> List[SingleTurnResult]:
    results: List[SingleTurnResult] = []
    total = len(questions)

    for i, q in enumerate(questions, 1):
        num = q["number"]
        category = q["category"]
        question = q["question"]
        session_id = f"00000000-0000-0000-0000-{str(num).zfill(12)}"

        print(f"  [{i:>3}/{total}] Q{num:>3}: {question[:70]}...", end="", flush=True)
        t0 = time.monotonic()
        try:
            data = _call_api(session_id, question, [])
            latency_ms = int((time.monotonic() - t0) * 1000)
            extracted = _extract_result(data)
            status = _classify_status(data)

            result = SingleTurnResult(
                number=num,
                category=category,
                question=question,
                status=status,
                path=extracted["path"],
                metric=extracted["metric"],
                confidence=extracted["confidence"],
                response_type=extracted["response_type"],
                answer=extracted["answer"],
                sql_used=extracted["sql_used"],
                row_count=extracted["row_count"],
                latency_ms=latency_ms,
                error="",
            )
            label = f"[{status}] {extracted['path']}"
            if extracted.get("metric"):
                label += f" ({extracted['metric']} {int((extracted['confidence'] or 0)*100)}%)"
            print(f"  {label}  {latency_ms}ms")

        except Exception as exc:
            latency_ms = int((time.monotonic() - t0) * 1000)
            result = SingleTurnResult(
                number=num, category=category, question=question,
                status=STATUS_ERROR, path="error",
                metric=None, confidence=None, response_type=None,
                answer="", sql_used="",
                row_count=None, latency_ms=latency_ms,
                error=str(exc),
            )
            print(f"  [ERROR] {exc}")

        results.append(result)

    return results


# ── Multi-turn runner ─────────────────────────────────────────────────────────

def run_multi_turn() -> List[MultiTurnResult]:
    from .multiturn_scenarios import SCENARIOS

    results: List[MultiTurnResult] = []

    for scenario in SCENARIOS:
        print(f"\n  Scenario {scenario.id}: {scenario.name}")
        session_id = f"00000000-0000-0000-8000-{str(scenario.id).zfill(12)}"
        history: List[Dict[str, str]] = []

        for turn_idx, turn in enumerate(scenario.turns, 1):
            print(f"    Turn {turn_idx}: {turn.question[:65]}...", end="", flush=True)
            t0 = time.monotonic()
            try:
                data = _call_api(session_id, turn.question, history)
                latency_ms = int((time.monotonic() - t0) * 1000)
                extracted = _extract_result(data)
                status = _classify_status(data)

                # Append to history for next turn
                history.append({"role": "user", "content": turn.question})
                history.append({"role": "assistant", "content": extracted["answer"]})

                result = MultiTurnResult(
                    scenario_id=scenario.id,
                    scenario_name=scenario.name,
                    scenario_description=scenario.description,
                    turn_number=turn_idx,
                    question=turn.question,
                    expected_intent=turn.expected_intent,
                    expected_path=turn.expected_path,
                    status=status,
                    path=extracted["path"],
                    metric=extracted["metric"],
                    confidence=extracted["confidence"],
                    response_type=extracted["response_type"],
                    answer=extracted["answer"],
                    sql_used=extracted["sql_used"],
                    latency_ms=latency_ms,
                    error="",
                    context_turns=turn_idx - 1,
                )
                path_match = (
                    "✓" if turn.expected_path is None or turn.expected_path == extracted["path"]
                    else "✗ expected " + (turn.expected_path or "?")
                )
                print(f"  [{status}] {path_match} {extracted['path']}  {latency_ms}ms")

            except Exception as exc:
                latency_ms = int((time.monotonic() - t0) * 1000)
                result = MultiTurnResult(
                    scenario_id=scenario.id,
                    scenario_name=scenario.name,
                    scenario_description=scenario.description,
                    turn_number=turn_idx,
                    question=turn.question,
                    expected_intent=turn.expected_intent,
                    expected_path=turn.expected_path,
                    status=STATUS_ERROR,
                    path="error",
                    metric=None, confidence=None, response_type=None,
                    answer="", sql_used="",
                    latency_ms=latency_ms,
                    error=str(exc),
                    context_turns=turn_idx - 1,
                )
                print(f"  [ERROR] {exc}")
                # Keep history intact for remaining turns so context persists
                history.append({"role": "user", "content": turn.question})
                history.append({"role": "assistant", "content": "[error]"})

            results.append(result)

    return results


# ── Excel report ──────────────────────────────────────────────────────────────

def _col_width(value: str, minimum: int = 10, maximum: int = 60) -> int:
    return min(maximum, max(minimum, len(str(value)) + 2))


def write_report(
    single: List[SingleTurnResult],
    multi: List[MultiTurnResult],
    output_path: Path,
    run_started: str,
) -> None:
    import openpyxl
    from openpyxl.styles import (
        Alignment, Border, Font, PatternFill, Side
    )
    from openpyxl.utils import get_column_letter

    wb = openpyxl.Workbook()

    # ── Colour palette ────────────────────────────────────────────────────────
    C_HEADER_BG   = "1E3A5F"   # dark navy
    C_HEADER_FG   = "FFFFFF"
    C_PASS        = "D4EDDA"   # green
    C_PARTIAL     = "FFF3CD"   # yellow
    C_ERROR       = "F8D7DA"   # red
    C_CERT        = "D1ECF1"   # teal — certified path
    C_SQL         = "EDE7F6"   # purple — text_to_sql path
    C_ALT_ROW     = "F7F9FC"   # light blue-grey for alternating rows
    C_SECTION_BG  = "2C5282"   # section header
    C_SECTION_FG  = "FFFFFF"

    thin_border = Border(
        left=Side(style="thin", color="D0D7DE"),
        right=Side(style="thin", color="D0D7DE"),
        top=Side(style="thin", color="D0D7DE"),
        bottom=Side(style="thin", color="D0D7DE"),
    )

    def header_fill(hex_colour: str) -> PatternFill:
        return PatternFill("solid", fgColor=hex_colour)

    def cell_fill(hex_colour: str) -> PatternFill:
        return PatternFill("solid", fgColor=hex_colour)

    def hdr_font(bold: bool = True) -> Font:
        return Font(name="Calibri", bold=bold, color=C_HEADER_FG, size=10)

    def body_font(bold: bool = False) -> Font:
        return Font(name="Calibri", bold=bold, size=10)

    def apply_header(ws, row: int, values: list) -> None:
        for col, val in enumerate(values, 1):
            c = ws.cell(row=row, column=col, value=val)
            c.font = hdr_font()
            c.fill = header_fill(C_HEADER_BG)
            c.border = thin_border
            c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    def status_fill(status: str) -> PatternFill:
        if status == STATUS_PASS:
            return cell_fill(C_PASS)
        if status == STATUS_PARTIAL:
            return cell_fill(C_PARTIAL)
        return cell_fill(C_ERROR)

    def path_fill(path: str) -> PatternFill:
        if path == "certified":
            return cell_fill(C_CERT)
        return cell_fill(C_SQL)

    # ── Sheet 1: All Questions ────────────────────────────────────────────────
    ws1 = wb.active
    ws1.title = "All Questions"
    ws1.freeze_panes = "A3"
    ws1.row_dimensions[1].height = 30
    ws1.row_dimensions[2].height = 22

    ws1.merge_cells("A1:O1")
    title_cell = ws1["A1"]
    title_cell.value = f"NEP Analytics — Test Suite Results   |   Run: {run_started}   |   {len(single)} questions"
    title_cell.font = Font(name="Calibri", bold=True, size=12, color=C_SECTION_FG)
    title_cell.fill = header_fill(C_SECTION_BG)
    title_cell.alignment = Alignment(horizontal="left", vertical="center")

    headers = [
        "#", "Category", "Question",
        "Status", "Path", "Certified Metric", "Confidence",
        "Response Type", "Answer", "SQL Used",
        "Rows", "Latency (ms)", "Error",
    ]
    apply_header(ws1, 2, headers)

    col_widths = [5, 30, 60, 10, 15, 25, 12, 15, 80, 60, 8, 12, 40]
    for i, w in enumerate(col_widths, 1):
        ws1.column_dimensions[get_column_letter(i)].width = w

    for r_idx, res in enumerate(single, 1):
        row = r_idx + 2
        alt = row % 2 == 0

        def wc(col, value, wrap=False, bold=False, fill=None, align="left"):
            c = ws1.cell(row=row, column=col, value=value)
            c.font = body_font(bold)
            c.border = thin_border
            c.alignment = Alignment(horizontal=align, vertical="top", wrap_text=wrap)
            if fill:
                c.fill = fill
            elif alt:
                c.fill = cell_fill(C_ALT_ROW)
            return c

        conf_str = f"{int(res.confidence * 100)}%" if res.confidence is not None else ""

        wc(1, res.number, align="center")
        wc(2, res.category, wrap=True)
        wc(3, res.question, wrap=True)
        wc(4, res.status, align="center", fill=status_fill(res.status), bold=True)
        wc(5, res.path, align="center", fill=path_fill(res.path))
        wc(6, res.metric or "", align="center")
        wc(7, conf_str, align="center",
           fill=cell_fill("C8E6C9") if res.confidence and res.confidence >= 0.8 else
                cell_fill("FFF9C4") if res.confidence and res.confidence >= 0.5 else None)
        wc(8, res.response_type or "", align="center")
        wc(9, res.answer, wrap=True)
        wc(10, res.sql_used, wrap=False)
        wc(11, res.row_count if res.row_count is not None else "", align="center")
        wc(12, res.latency_ms, align="center")
        wc(13, res.error, wrap=True)

        ws1.row_dimensions[row].height = 45 if len(res.question) > 80 else 30

    # ── Sheet 2: Multi-Turn Scenarios ─────────────────────────────────────────
    ws2 = wb.create_sheet("Multi-Turn Scenarios")
    ws2.freeze_panes = "A3"
    ws2.row_dimensions[1].height = 30
    ws2.row_dimensions[2].height = 22

    ws2.merge_cells("A1:N1")
    t2 = ws2["A1"]
    t2.value = f"Multi-Turn Conversation Tests   |   {len(SCENARIOS_COUNT(multi))} scenarios, {len(multi)} turns total"
    t2.font = Font(name="Calibri", bold=True, size=12, color=C_SECTION_FG)
    t2.fill = header_fill(C_SECTION_BG)
    t2.alignment = Alignment(horizontal="left", vertical="center")

    mt_headers = [
        "Scenario", "Scenario Name", "Turn #", "Prior Turns",
        "Question", "Expected Intent", "Expected Path",
        "Status", "Actual Path", "Metric", "Confidence",
        "Response Type", "Answer", "Error",
    ]
    apply_header(ws2, 2, mt_headers)

    mt_widths = [8, 28, 7, 10, 55, 50, 15, 10, 15, 22, 12, 15, 80, 40]
    for i, w in enumerate(mt_widths, 1):
        ws2.column_dimensions[get_column_letter(i)].width = w

    prev_scenario = None
    for r_idx, res in enumerate(multi, 1):
        row = r_idx + 2
        new_scenario = res.scenario_id != prev_scenario
        prev_scenario = res.scenario_id

        def wm(col, value, wrap=False, bold=False, fill=None, align="left"):
            c = ws2.cell(row=row, column=col, value=value)
            c.font = body_font(bold)
            c.border = thin_border
            c.alignment = Alignment(horizontal=align, vertical="top", wrap_text=wrap)
            if fill:
                c.fill = fill
            elif new_scenario and r_idx % 2 == 0:
                c.fill = cell_fill("EEF2FF")
            return c

        path_match_str = ""
        if res.expected_path and res.path != "error":
            path_match_str = "✓" if res.expected_path == res.path else f"✗ (exp: {res.expected_path})"

        conf_str = f"{int(res.confidence * 100)}%" if res.confidence is not None else ""

        wm(1, res.scenario_id, align="center", bold=new_scenario)
        wm(2, res.scenario_name, wrap=True, bold=new_scenario)
        wm(3, res.turn_number, align="center")
        wm(4, res.context_turns, align="center")
        wm(5, res.question, wrap=True)
        wm(6, res.expected_intent, wrap=True)
        wm(7, res.expected_path or "any", align="center")
        wm(8, res.status, align="center", fill=status_fill(res.status), bold=True)
        wm(9, f"{res.path} {path_match_str}", align="center", fill=path_fill(res.path))
        wm(10, res.metric or "", align="center")
        wm(11, conf_str, align="center")
        wm(12, res.response_type or "", align="center")
        wm(13, res.answer, wrap=True)
        wm(14, res.error, wrap=True)

        ws2.row_dimensions[row].height = 50

    # ── Sheet 3: Summary ──────────────────────────────────────────────────────
    ws3 = wb.create_sheet("Summary")

    def sec(ws, row, label):
        ws.merge_cells(f"A{row}:D{row}")
        c = ws.cell(row=row, column=1, value=label)
        c.font = Font(name="Calibri", bold=True, size=11, color=C_SECTION_FG)
        c.fill = header_fill(C_SECTION_BG)
        c.alignment = Alignment(horizontal="left", vertical="center")
        ws.row_dimensions[row].height = 22

    def kv(ws, row, key, value, value_fill=None):
        kc = ws.cell(row=row, column=1, value=key)
        kc.font = body_font(bold=True)
        kc.border = thin_border
        kc.fill = cell_fill("EBF5FF")
        vc = ws.cell(row=row, column=2, value=value)
        vc.font = body_font()
        vc.border = thin_border
        if value_fill:
            vc.fill = cell_fill(value_fill)

    ws3.column_dimensions["A"].width = 38
    ws3.column_dimensions["B"].width = 18
    ws3.column_dimensions["C"].width = 30
    ws3.column_dimensions["D"].width = 18

    total = len(single)
    passed = sum(1 for r in single if r.status == STATUS_PASS)
    partial = sum(1 for r in single if r.status == STATUS_PARTIAL)
    errors = sum(1 for r in single if r.status == STATUS_ERROR)
    certified_count = sum(1 for r in single if r.path == "certified")
    text_sql_count = sum(1 for r in single if r.path == "text_to_sql")
    cert_confidences = [r.confidence for r in single if r.confidence is not None]
    avg_confidence = sum(cert_confidences) / len(cert_confidences) if cert_confidences else 0
    avg_latency = sum(r.latency_ms for r in single) / total if total else 0
    p95_latency = sorted(r.latency_ms for r in single)[int(total * 0.95)] if total > 20 else max((r.latency_ms for r in single), default=0)

    r = 1
    sec(ws3, r, "Single-Turn Test Summary"); r += 1
    kv(ws3, r, "Total Questions Tested", total); r += 1
    kv(ws3, r, "PASS", f"{passed}  ({round(passed/total*100,1)}%)", "D4EDDA"); r += 1
    kv(ws3, r, "PARTIAL (DB/SQL error in answer)", f"{partial}  ({round(partial/total*100,1)}%)", "FFF3CD"); r += 1
    kv(ws3, r, "ERROR (HTTP/exception)", f"{errors}  ({round(errors/total*100,1)}%)", "F8D7DA"); r += 1
    kv(ws3, r, "Certified metric path", f"{certified_count}  ({round(certified_count/total*100,1)}%)"); r += 1
    kv(ws3, r, "Text-to-SQL path", f"{text_sql_count}  ({round(text_sql_count/total*100,1)}%)"); r += 1
    kv(ws3, r, "Avg confidence (certified answers)", f"{round(avg_confidence*100,1)}%"); r += 1
    kv(ws3, r, "Avg response latency", f"{round(avg_latency/1000,1)}s"); r += 1
    kv(ws3, r, "P95 response latency", f"{round(p95_latency/1000,1)}s"); r += 1

    r += 1
    sec(ws3, r, "Results by Category"); r += 1
    ws3.cell(row=r, column=1, value="Category").font = body_font(bold=True)
    ws3.cell(row=r, column=2, value="Total").font = body_font(bold=True)
    ws3.cell(row=r, column=3, value="PASS").font = body_font(bold=True)
    ws3.cell(row=r, column=4, value="PARTIAL/ERROR").font = body_font(bold=True)
    for c in range(1, 5):
        ws3.cell(row=r, column=c).fill = header_fill(C_HEADER_BG)
        ws3.cell(row=r, column=c).font = hdr_font()
        ws3.cell(row=r, column=c).border = thin_border
    r += 1

    from collections import defaultdict
    cat_stats: Dict[str, Dict] = defaultdict(lambda: {"total": 0, "pass": 0, "fail": 0})
    for res in single:
        cat_stats[res.category]["total"] += 1
        if res.status == STATUS_PASS:
            cat_stats[res.category]["pass"] += 1
        else:
            cat_stats[res.category]["fail"] += 1

    for cat, stats in sorted(cat_stats.items()):
        ws3.cell(row=r, column=1, value=cat).border = thin_border
        ws3.cell(row=r, column=2, value=stats["total"]).border = thin_border
        ws3.cell(row=r, column=3, value=stats["pass"]).border = thin_border
        pct_pass = round(stats["pass"] / stats["total"] * 100, 1) if stats["total"] else 0
        ws3.cell(row=r, column=3, value=f"{stats['pass']} ({pct_pass}%)").border = thin_border
        ws3.cell(row=r, column=4, value=stats["fail"]).border = thin_border
        ws3.cell(row=r, column=1).fill = cell_fill("F7F9FC")
        r += 1

    r += 1
    sec(ws3, r, "Multi-Turn Scenario Summary"); r += 1
    mt_pass = sum(1 for x in multi if x.status == STATUS_PASS)
    mt_total = len(multi)
    kv(ws3, r, "Scenarios run", len(set(x.scenario_id for x in multi))); r += 1
    kv(ws3, r, "Total turns", mt_total); r += 1
    kv(ws3, r, "Turns PASS", f"{mt_pass}  ({round(mt_pass/mt_total*100,1) if mt_total else 0}%)", "D4EDDA"); r += 1
    mt_path_match = sum(
        1 for x in multi
        if x.expected_path and x.path == x.expected_path
    )
    mt_path_expected = sum(1 for x in multi if x.expected_path)
    kv(ws3, r, "Path routing correct (where expected set)",
       f"{mt_path_match}/{mt_path_expected}  ({round(mt_path_match/mt_path_expected*100,1) if mt_path_expected else 0}%)"); r += 1

    r += 1
    sec(ws3, r, "Errors Log"); r += 1
    ws3.cell(row=r, column=1, value="#").font = body_font(bold=True)
    ws3.cell(row=r, column=2, value="Question").font = body_font(bold=True)
    ws3.cell(row=r, column=3, value="Error").font = body_font(bold=True)
    ws3.merge_cells(f"B{r}:B{r}")
    ws3.column_dimensions["B"].width = 55
    r += 1

    err_items = [(res.number, res.question, res.error) for res in single if res.status != STATUS_PASS]
    if not err_items:
        ws3.cell(row=r, column=1, value="No errors").font = body_font()
        r += 1
    else:
        for num, q, err in err_items:
            ws3.cell(row=r, column=1, value=num).border = thin_border
            qc = ws3.cell(row=r, column=2, value=q[:80])
            qc.border = thin_border
            qc.alignment = Alignment(wrap_text=True)
            ec = ws3.cell(row=r, column=3, value=err[:120])
            ec.border = thin_border
            ec.alignment = Alignment(wrap_text=True)
            ws3.row_dimensions[r].height = 30
            r += 1

    wb.save(output_path)
    print(f"\n  Report saved → {output_path}")


def SCENARIOS_COUNT(multi_results):
    return set(r.scenario_id for r in multi_results)


# ── Question loader ───────────────────────────────────────────────────────────

def load_questions(
    category_filter: Optional[str] = None,
    limit: Optional[int] = None,
) -> List[Dict[str, Any]]:
    import openpyxl
    wb = openpyxl.load_workbook(INPUT_XLSX)
    ws = wb["Test Questions"]
    questions = []
    for row in ws.iter_rows(min_row=2, values_only=True):
        num, category, question = row[0], row[1], row[2]
        if not num or not question:
            continue
        if category_filter and category_filter.lower() not in category.lower():
            continue
        questions.append({"number": int(num), "category": category, "question": question})
    if limit:
        questions = questions[:limit]
    return questions


# ── Main entry point ──────────────────────────────────────────────────────────

def run(
    category: Optional[str] = None,
    limit: Optional[int] = None,
    skip_multiturn: bool = False,
    output_path: Optional[Path] = None,
) -> None:
    output_path = output_path or DEFAULT_OUTPUT
    run_started = datetime.now().strftime("%Y-%m-%d %H:%M")

    print(f"\n{'='*64}")
    print(f"  NEP Analytics Test Suite")
    print(f"  Started: {run_started}")
    print(f"  Adapter: {BASE_URL}")
    print(f"{'='*64}\n")

    questions = load_questions(category_filter=category, limit=limit)
    print(f"  Single-turn questions: {len(questions)}")
    if not questions:
        print("  No questions loaded. Check the input file path.")
        sys.exit(1)

    print(f"\n  Running {len(questions)} single-turn questions...")
    single_results = run_single_turn(questions)

    multi_results: List[MultiTurnResult] = []
    if not skip_multiturn:
        from .multiturn_scenarios import SCENARIOS
        print(f"\n  Running {len(SCENARIOS)} multi-turn scenarios ({sum(len(s.turns) for s in SCENARIOS)} turns)...")
        multi_results = run_multi_turn()

    print(f"\n  Writing Excel report...")
    write_report(single_results, multi_results, output_path, run_started)

    # Print final summary
    total = len(single_results)
    passed = sum(1 for r in single_results if r.status == STATUS_PASS)
    partial = sum(1 for r in single_results if r.status == STATUS_PARTIAL)
    errors = sum(1 for r in single_results if r.status == STATUS_ERROR)
    cert = sum(1 for r in single_results if r.path == "certified")

    print(f"\n{'='*64}")
    print(f"  Single-turn: {passed}/{total} PASS  |  {partial} PARTIAL  |  {errors} ERROR")
    print(f"  Certified path: {cert}/{total} ({round(cert/total*100,1)}%)")
    if multi_results:
        mt_pass = sum(1 for r in multi_results if r.status == STATUS_PASS)
        print(f"  Multi-turn: {mt_pass}/{len(multi_results)} turns PASS")
    print(f"{'='*64}\n")
