"""
NEP Analytics — Live DB Fact Extraction
Runs all 24 questions from nep_query_test_suite.xlsx against the live DB
and reports ACTUAL results (SQL + answer). No expected-value comparison.

Usage (from repo root):
    API_KEY=<key> python nep_analytics/tests/run_suite_live.py
"""
from __future__ import annotations

import asyncio
import os
import sys
import time
import uuid
from pathlib import Path

import httpx
import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

HERE      = Path(__file__).resolve().parent
REPO_ROOT = HERE.parent.parent
sys.path.insert(0, str(REPO_ROOT))

API_BASE        = os.getenv("API_BASE_URL", "http://localhost:8002")
API_KEY         = os.getenv("API_KEY", "nep-analytics-internal-2026")
HEADERS         = {"Content-Type": "application/json", "x-api-key": API_KEY}
REQUEST_TIMEOUT = 90.0
MAX_CONCURRENT  = 4

SUITE_FILE  = HERE / "nep_query_test_suite.xlsx"
REPORT_FILE = HERE / "NEP_Suite_Live_Results.xlsx"

# ── Palette ───────────────────────────────────────────────────────────────────
C_DARK  = "002C5282"
C_MID   = "001E3A5F"
C_LIGHT = "00EBF5FF"
C_ALT   = "00F7F9FC"
C_WHITE = "00FFFFFF"
C_OK    = "00D4EDDA"
C_WARN  = "00FFF3CD"
C_ERR   = "00F8D7DA"
C_DNA   = "00E9ECEF"

BEHAVIOR_LABEL = {
    "rows":            "Happy Path",
    "grain":           "Grain Trap",
    "normalize":       "Normalization",
    "correctly_empty": "Correctly Empty",
    "unanswerable":    "Unanswerable",
}


# ── Load questions only ───────────────────────────────────────────────────────

def load_questions() -> list[dict]:
    wb = openpyxl.load_workbook(SUITE_FILE)
    ws = wb["Test Cases"]
    cases = []
    for r in range(5, ws.max_row + 1):
        row = [ws.cell(r, c).value for c in range(1, ws.max_column + 1)]
        if not row[0]:
            continue
        cases.append({
            "id":       str(row[0]).strip(),
            "category": str(row[1]).strip() if row[1] else "",
            "question": str(row[2]).strip() if row[2] else "",
            "behavior": str(row[4]).strip() if row[4] else "",
        })
    return cases


# ── API call ──────────────────────────────────────────────────────────────────

async def ask_api(client: httpx.AsyncClient, sem: asyncio.Semaphore,
                  case: dict) -> dict:
    payload = {"question": case["question"], "session_id": str(uuid.uuid4())}
    t0 = time.perf_counter()
    answer, sql_used, response_type, row_count, error, http_status = \
        "", "", "text", None, None, 200

    try:
        async with sem:
            r = await client.post(
                f"{API_BASE}/api/chat",
                json=payload,
                headers=HEADERS,
                timeout=REQUEST_TIMEOUT,
            )
            http_status = r.status_code
            if http_status == 200:
                data          = r.json()
                answer        = data.get("answer", "")
                sql_used      = data.get("sql_used", "") or ""
                response_type = data.get("response_type", "text")
                # Extract row count from table/chart if present
                tbl = data.get("table_data") or {}
                if isinstance(tbl.get("rows"), list):
                    row_count = len(tbl["rows"])
                chart = data.get("chart_config") or {}
                if isinstance(chart.get("labels"), list):
                    row_count = len(chart["labels"])
            else:
                error = f"HTTP {http_status}: {r.text[:200]}"
    except Exception as exc:
        error       = str(exc)[:300]
        http_status = 0

    latency_ms = int((time.perf_counter() - t0) * 1000)

    # Simple status classification
    if error or http_status >= 500:
        status = "ERROR"
    elif not sql_used:
        status = "CLARIFICATION"
    else:
        ans_lower = (answer or "").lower()
        no_data_phrases = [
            "no data", "no results", "no matching", "no records",
            "0 results", "zero results", "not find", "no users found",
            "no events found", "no mentors", "does not exist",
            "no information available", "no entries", "no rows",
        ]
        if any(p in ans_lower for p in no_data_phrases):
            status = "ZERO_RESULTS"
        elif row_count == 0:
            status = "ZERO_RESULTS"
        else:
            status = "HAS_RESULTS"

    return {
        **case,
        "http_status":   http_status,
        "latency_ms":    latency_ms,
        "status":        status,
        "answer":        answer,
        "sql_used":      sql_used,
        "response_type": response_type,
        "row_count":     row_count,
        "error":         error,
    }


# ── Excel report ──────────────────────────────────────────────────────────────

def _thin():
    s = Side(style="thin", color="FFAAAAAA")
    return Border(left=s, right=s, top=s, bottom=s)

def _fill(h):
    return PatternFill("solid", fgColor=h)

def _font(bold=False, color="FF000000", size=10):
    return Font(bold=bold, color=color, size=size)

def _align(h="left", v="top", wrap=True):
    return Alignment(horizontal=h, vertical=v, wrap_text=wrap)


STATUS_FILL = {
    "HAS_RESULTS":   C_OK,
    "ZERO_RESULTS":  C_WARN,
    "CLARIFICATION": C_LIGHT,
    "ERROR":         C_ERR,
}


def write_excel(results: list[dict]) -> Path:
    wb  = openpyxl.Workbook()
    ws  = wb.active
    ws.title = "Live Results"

    col_defs = [
        ("ID",              7),
        ("Category",        18),
        ("Question",        48),
        ("Status",          14),
        ("Latency (ms)",    13),
        ("Response Type",   14),
        ("Row Count",       11),
        ("Answer",          55),
        ("SQL Used",        70),
    ]

    # Title
    ws.merge_cells(f"A1:{chr(64+len(col_defs))}1")
    ws["A1"] = "NEP Analytics — Live DB Fact Extraction"
    ws["A1"].font      = Font(bold=True, color="FFFFFFFF", size=13)
    ws["A1"].fill      = _fill(C_DARK)
    ws["A1"].alignment = _align("center", "center")
    ws.row_dimensions[1].height = 22

    # Headers
    for ci, (h, w) in enumerate(col_defs, 1):
        cell = ws.cell(2, ci, h)
        cell.font      = _font(bold=True, color="FFFFFFFF")
        cell.fill      = _fill(C_MID)
        cell.alignment = _align("center", "center", wrap=False)
        cell.border    = _thin()
        ws.column_dimensions[cell.column_letter].width = w
    ws.row_dimensions[2].height = 18
    ws.freeze_panes = "A3"

    # Group rows by behavior category
    by_behavior: dict[str, list[dict]] = {}
    for r in results:
        lbl = BEHAVIOR_LABEL.get(r["behavior"], r["category"])
        by_behavior.setdefault(lbl, []).append(r)

    row_idx = 3
    for behavior_label, rows in by_behavior.items():
        # Category separator
        ws.merge_cells(f"A{row_idx}:{chr(64+len(col_defs))}{row_idx}")
        cat_cell = ws.cell(row_idx, 1, f"  {behavior_label}")
        cat_cell.font      = _font(bold=True, color="FFFFFFFF", size=10)
        cat_cell.fill      = _fill(C_MID)
        cat_cell.alignment = _align("left", "center", wrap=False)
        ws.row_dimensions[row_idx].height = 16
        row_idx += 1

        for res in rows:
            status      = res["status"]
            row_fill    = STATUS_FILL.get(status, C_WHITE)
            alt_fill    = C_ALT if row_idx % 2 == 0 else C_WHITE
            rc          = res.get("row_count")
            rc_display  = str(rc) if rc is not None else "—"

            vals = [
                res["id"],
                BEHAVIOR_LABEL.get(res["behavior"], res["category"]),
                res["question"],
                status,
                res["latency_ms"],
                res["response_type"],
                rc_display,
                res["answer"] or res.get("error", ""),
                res["sql_used"],
            ]
            for ci, val in enumerate(vals, 1):
                cell = ws.cell(row_idx, ci, val)
                cell.border    = _thin()
                cell.alignment = _align(wrap=True)
                if ci == 4:  # Status column
                    cell.fill = _fill(row_fill)
                    cell.font = _font(bold=True)
                    cell.alignment = _align("center", "center", wrap=False)
                else:
                    cell.fill = _fill(alt_fill)
                    cell.font = _font()
            ws.row_dimensions[row_idx].height = 80
            row_idx += 1

    # ── Sheet 2: Summary ─────────────────────────────────────────────────────
    ws2 = wb.create_sheet("Summary")

    ws2.merge_cells("A1:E1")
    ws2["A1"] = "Live DB Fact Extraction — Summary"
    ws2["A1"].font      = Font(bold=True, color="FFFFFFFF", size=13)
    ws2["A1"].fill      = _fill(C_DARK)
    ws2["A1"].alignment = _align("center", "center")
    ws2.row_dimensions[1].height = 22

    headers = ["Category", "Total", "Has Results", "Zero Results", "Clarification / Error"]
    for ci, h in enumerate(headers, 1):
        cell = ws2.cell(3, ci, h)
        cell.font      = _font(bold=True, color="FFFFFFFF")
        cell.fill      = _fill(C_MID)
        cell.border    = _thin()
        cell.alignment = _align("center", "center", wrap=False)
        ws2.column_dimensions[cell.column_letter].width = 24

    ri = 4
    for behavior_label, rows in by_behavior.items():
        total    = len(rows)
        has_res  = sum(1 for r in rows if r["status"] == "HAS_RESULTS")
        zero_res = sum(1 for r in rows if r["status"] == "ZERO_RESULTS")
        other    = total - has_res - zero_res
        row_vals = [behavior_label, total, has_res, zero_res, other]
        for ci, val in enumerate(row_vals, 1):
            cell = ws2.cell(ri, ci, val)
            cell.border    = _thin()
            cell.alignment = _align("center", "center", wrap=False)
            cell.fill      = _fill(C_ALT if ri % 2 == 0 else C_WHITE)
        ri += 1

    # Totals row
    all_total    = len(results)
    all_has_res  = sum(1 for r in results if r["status"] == "HAS_RESULTS")
    all_zero_res = sum(1 for r in results if r["status"] == "ZERO_RESULTS")
    all_other    = all_total - all_has_res - all_zero_res
    for ci, val in enumerate(["TOTAL", all_total, all_has_res, all_zero_res, all_other], 1):
        cell = ws2.cell(ri, ci, val)
        cell.font   = _font(bold=True)
        cell.fill   = _fill(C_MID)
        cell.border = _thin()
        cell.alignment = _align("center", "center", wrap=False)
        if ci == 1:
            cell.font = _font(bold=True, color="FFFFFFFF")

    wb.save(REPORT_FILE)
    return REPORT_FILE


# ── Main ──────────────────────────────────────────────────────────────────────

async def main():
    cases = load_questions()
    print(f"\nNEP Analytics — Live DB Fact Extraction")
    print(f"  Questions : {len(cases)}")
    print(f"  API       : {API_BASE}\n")

    sem = asyncio.Semaphore(MAX_CONCURRENT)
    async with httpx.AsyncClient() as client:
        tasks   = [ask_api(client, sem, c) for c in cases]
        results = []
        for coro in asyncio.as_completed(tasks):
            res  = await coro
            icon = {"HAS_RESULTS": "✓", "ZERO_RESULTS": "◌",
                    "CLARIFICATION": "?", "ERROR": "✗"}.get(res["status"], " ")
            rc   = f"rows={res['row_count']}" if res["row_count"] is not None else ""
            print(f"  [{icon}] {res['id']:4s}  {res['latency_ms']:5d}ms  "
                  f"{res['status']:<14}  {rc:10}  {res['question'][:55]}")
            results.append(res)

    results.sort(key=lambda r: r["id"])
    out = write_excel(results)

    # Console summary
    by_status: dict[str, int] = {}
    for r in results:
        by_status[r["status"]] = by_status.get(r["status"], 0) + 1

    print(f"\n{'='*60}")
    print(f"  24 questions — Live DB Results")
    print(f"{'='*60}")
    for s, cnt in sorted(by_status.items()):
        icon = {"HAS_RESULTS": "✓", "ZERO_RESULTS": "◌",
                "CLARIFICATION": "?", "ERROR": "✗"}.get(s, " ")
        print(f"  [{icon}] {s:<18}: {cnt}")
    print(f"\n  By category:")
    by_beh: dict[str, dict] = {}
    for r in results:
        lbl = BEHAVIOR_LABEL.get(r["behavior"], r["behavior"])
        bc  = by_beh.setdefault(lbl, {"has": 0, "zero": 0, "other": 0, "total": 0})
        bc["total"] += 1
        if r["status"] == "HAS_RESULTS":
            bc["has"] += 1
        elif r["status"] == "ZERO_RESULTS":
            bc["zero"] += 1
        else:
            bc["other"] += 1
    for lbl, bc in by_beh.items():
        print(f"    {lbl:<22}  has={bc['has']}  zero={bc['zero']}  other={bc['other']}")
    print(f"\n  Report → {out}")
    print(f"{'='*60}\n")


if __name__ == "__main__":
    asyncio.run(main())
