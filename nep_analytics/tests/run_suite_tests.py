"""
NEP Analytics — Production Simulation Test Suite Runner
Executes all 24 test cases from nep_query_test_suite.xlsx against
the nep_analytics API and produces a scored Excel report.

Usage (from repo root):
    API_KEY=<key> python nep_analytics/tests/run_suite_tests.py
"""
from __future__ import annotations

import asyncio
import json
import os
import sys
import time
import uuid
from pathlib import Path

import httpx
import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

HERE = Path(__file__).resolve().parent
REPO_ROOT = HERE.parent.parent
sys.path.insert(0, str(REPO_ROOT))

from nep_analytics.core.config import get_anthropic  # noqa: E402

# ── Config ────────────────────────────────────────────────────────────────────
API_BASE = os.getenv("API_BASE_URL", "http://localhost:8002")
API_KEY  = os.getenv("API_KEY", "nep-analytics-internal-2026")
HEADERS  = {"Content-Type": "application/json", "x-api-key": API_KEY}
REQUEST_TIMEOUT = 90.0
MAX_CONCURRENT  = 4
JUDGE_MODEL     = "claude-opus-5"

SUITE_FILE  = HERE / "nep_query_test_suite.xlsx"
REPORT_FILE = HERE / "NEP_Suite_Test_Report.xlsx"

# ── Colour palette ────────────────────────────────────────────────────────────
C_DARK   = "002C5282"
C_MID    = "001E3A5F"
C_LIGHT  = "00EBF5FF"
C_ALT    = "00F7F9FC"
C_PASS   = "00D4EDDA"
C_WARN   = "00FFF3CD"
C_FAIL   = "00F8D7DA"
C_INFO   = "00D1ECF1"
C_GREY   = "00E9ECEF"
C_WHITE  = "00FFFFFF"

VERDICT_FILL = {
    "PASS":    C_PASS,
    "PARTIAL": C_WARN,
    "FAIL":    C_FAIL,
    "SKIP":    C_GREY,
}

BEHAVIOR_LABEL = {
    "rows":            "Happy Path",
    "grain":           "Grain Trap",
    "normalize":       "Normalization",
    "correctly_empty": "Correctly Empty",
    "unanswerable":    "Unanswerable",
}


# ── Read test cases ───────────────────────────────────────────────────────────

def load_test_cases() -> list[dict]:
    wb = openpyxl.load_workbook(SUITE_FILE)
    ws = wb["Test Cases"]
    cases = []
    for r in range(5, ws.max_row + 1):
        row = [ws.cell(r, c).value for c in range(1, ws.max_column + 1)]
        if not row[0]:
            continue
        cases.append({
            "id":            str(row[0]).strip(),
            "category":      str(row[1]).strip() if row[1] else "",
            "question":      str(row[2]).strip() if row[2] else "",
            "target_tables": str(row[3]).strip() if row[3] else "",
            "behavior":      str(row[4]).strip() if row[4] else "",
            "expected":      str(row[5]).strip() if row[5] else "",
            "wrong_if":      str(row[6]).strip() if row[6] else "",
            "gotcha":        str(row[7]).strip() if row[7] else "",
        })
    return cases


# ── API call ──────────────────────────────────────────────────────────────────

async def ask_api(client: httpx.AsyncClient, sem: asyncio.Semaphore,
                  case: dict) -> dict:
    session_id = str(uuid.uuid4())
    payload = {
        "question":   case["question"],
        "session_id": session_id,
    }
    t0 = time.perf_counter()
    status, answer, sql_used, response_type, error = 200, "", "", "text", None
    try:
        async with sem:
            r = await client.post(
                f"{API_BASE}/api/chat",
                json=payload,
                headers=HEADERS,
                timeout=REQUEST_TIMEOUT,
            )
            status = r.status_code
            if status == 200:
                data = r.json()
                answer        = data.get("answer", "")
                sql_used      = data.get("sql_used", "")
                response_type = data.get("response_type", "text")
            else:
                error = f"HTTP {status}: {r.text[:200]}"
    except Exception as exc:
        error = str(exc)[:300]
        status = 0

    latency_ms = int((time.perf_counter() - t0) * 1000)
    return {
        **case,
        "http_status":   status,
        "latency_ms":    latency_ms,
        "answer":        answer,
        "sql_used":      sql_used,
        "response_type": response_type,
        "error":         error,
    }


# ── LLM Judge ────────────────────────────────────────────────────────────────

JUDGE_SYSTEM = """You are an expert evaluator for an analytics chatbot.
Given a natural-language question, expected behavior, expected value, and the
chatbot's actual response, judge whether the system handled the question correctly.

Return ONLY valid JSON — no markdown, no extra prose:
{
  "verdict": "PASS" | "PARTIAL" | "FAIL",
  "confidence_pct": <integer 0-100>,
  "reasoning": "<2-4 sentence explanation>"
}

Verdict rules:
  PASS    — The response fully satisfies the question and matches the expected value/behavior.
  PARTIAL — The response is on the right track but is incomplete or has a minor error.
  FAIL    — The response is wrong, uses wrong SQL, misses the key insight, or ignores the
            expected empty/unanswerable classification.

Behavior types and their pass criteria:
  rows            — SQL is correct AND returned values match the expected numbers.
  grain           — Used COUNT(DISTINCT ...) correctly; did NOT return the raw row count.
  normalize       — Used the correct enum values (camelCase, exact case) in the SQL.
  correctly_empty — Returned 0 results AND gave a clear explanation (did NOT error or panic).
  unanswerable    — Declined/explained the limitation clearly rather than generating wrong SQL.
"""


async def judge_one(client_sem_tuple, result: dict) -> dict:
    client, sem, anthropic = client_sem_tuple
    behavior = result["behavior"]
    expected = result["expected"]
    wrong_if = result["wrong_if"]
    gotcha   = result["gotcha"]

    user_msg = f"""Question: {result['question']}
Expected behavior: {BEHAVIOR_LABEL.get(behavior, behavior)}
Expected value/reason: {expected}
Wrong if: {wrong_if or 'n/a'}
Gotcha: {gotcha or 'n/a'}

--- ACTUAL RESPONSE ---
SQL used: {(result['sql_used'] or 'none')[:800]}
Answer: {(result['answer'] or 'none')[:1000]}
HTTP status: {result['http_status']}
Error: {result['error'] or 'none'}
"""

    try:
        async with sem:
            resp = await asyncio.get_event_loop().run_in_executor(
                None,
                lambda: anthropic.messages.create(
                    model=JUDGE_MODEL,
                    max_tokens=1500,
                    system=JUDGE_SYSTEM,
                    messages=[{"role": "user", "content": user_msg}],
                ),
            )
        raw = next((b.text for b in resp.content if b.type == "text"), "").strip()
        verdict_data = json.loads(raw)
    except Exception as exc:
        verdict_data = {
            "verdict":        "SKIP",
            "confidence_pct": 0,
            "reasoning":      f"Judge error: {exc}",
        }

    return {**result, "judge": verdict_data}


# ── Excel report ──────────────────────────────────────────────────────────────

def _thin() -> Border:
    s = Side(style="thin", color="FFAAAAAA")
    return Border(left=s, right=s, top=s, bottom=s)

def _fill(hex_color: str) -> PatternFill:
    return PatternFill("solid", fgColor=hex_color)

def _font(bold=False, color="FF000000", size=10):
    return Font(bold=bold, color=color, size=size)

def _align(h="left", v="top", wrap=True):
    return Alignment(horizontal=h, vertical=v, wrap_text=wrap)


def write_excel(results: list[dict], summary: dict) -> Path:
    wb = openpyxl.Workbook()

    # ── Sheet 1: Results ─────────────────────────────────────────────────────
    ws = wb.active
    ws.title = "Results"

    col_widths = [6, 18, 52, 14, 24, 26, 26, 12, 10, 60]
    col_headers = [
        "ID", "Category", "Query", "Behavior", "Expected Value",
        "SQL Used (excerpt)", "Answer (excerpt)", "Verdict", "Confidence", "Reasoning",
    ]

    # Title row
    ws.merge_cells("A1:J1")
    ws["A1"] = "NEP Analytics — Production Simulation Test Report"
    ws["A1"].font      = Font(bold=True, color="FFFFFFFF", size=13)
    ws["A1"].fill      = _fill(C_DARK)
    ws["A1"].alignment = _align("center", "center")
    ws.row_dimensions[1].height = 22

    # Header row
    for ci, (h, w) in enumerate(zip(col_headers, col_widths), start=1):
        cell = ws.cell(2, ci, h)
        cell.font      = _font(bold=True, color="FFFFFFFF", size=10)
        cell.fill      = _fill(C_MID)
        cell.alignment = _align("center", "center", wrap=False)
        cell.border    = _thin()
        ws.column_dimensions[cell.column_letter].width = w
    ws.row_dimensions[2].height = 18
    ws.freeze_panes = "A3"

    for ri, res in enumerate(results, start=3):
        j       = res.get("judge", {})
        verdict = j.get("verdict", "SKIP")
        row_fill_hex = VERDICT_FILL.get(verdict, C_WHITE)
        alt_fill_hex = C_ALT if ri % 2 == 0 else C_WHITE

        sql_excerpt = (res.get("sql_used") or "")[:300]
        ans_excerpt = (res.get("answer")   or "")[:300]
        reasoning   = j.get("reasoning", "")

        vals = [
            res["id"],
            BEHAVIOR_LABEL.get(res["behavior"], res["category"]),
            res["question"],
            BEHAVIOR_LABEL.get(res["behavior"], res["behavior"]),
            res["expected"],
            sql_excerpt,
            ans_excerpt,
            verdict,
            f"{j.get('confidence_pct', '')}%",
            reasoning,
        ]
        for ci, val in enumerate(vals, start=1):
            cell = ws.cell(ri, ci, val)
            cell.border    = _thin()
            cell.alignment = _align(wrap=True)
            if ci == 8:  # Verdict column
                cell.fill = _fill(row_fill_hex)
                cell.font = _font(bold=True)
                cell.alignment = _align("center", "center", wrap=False)
            else:
                cell.fill = _fill(alt_fill_hex)
                cell.font = _font()
        ws.row_dimensions[ri].height = 60

    # ── Sheet 2: Summary ─────────────────────────────────────────────────────
    ws2 = wb.create_sheet("Summary")

    ws2.merge_cells("A1:D1")
    ws2["A1"] = "Test Suite Summary"
    ws2["A1"].font      = Font(bold=True, color="FFFFFFFF", size=13)
    ws2["A1"].fill      = _fill(C_DARK)
    ws2["A1"].alignment = _align("center", "center")
    ws2.row_dimensions[1].height = 22

    # Overall counts
    ws2["A3"] = "Overall"
    ws2["A3"].font = _font(bold=True, size=11)

    overall_headers = ["Verdict", "Count", "%"]
    for ci, h in enumerate(overall_headers, start=1):
        cell = ws2.cell(4, ci, h)
        cell.font  = _font(bold=True, color="FFFFFFFF")
        cell.fill  = _fill(C_MID)
        cell.alignment = _align("center", "center", wrap=False)
        cell.border = _thin()

    total = summary["total"]
    for ri, (verdict, cnt) in enumerate(summary["overall"].items(), start=5):
        pct = f"{cnt/total*100:.0f}%" if total else "0%"
        for ci, val in enumerate([verdict, cnt, pct], start=1):
            cell = ws2.cell(ri, ci, val)
            cell.fill   = _fill(VERDICT_FILL.get(verdict, C_WHITE))
            cell.border = _thin()
            cell.alignment = _align("center", "center", wrap=False)
            if ci == 1:
                cell.font = _font(bold=True)

    # By category
    row_off = 5 + len(summary["overall"]) + 2
    ws2.cell(row_off, 1, "By Behavior Category").font = _font(bold=True, size=11)
    row_off += 1

    cat_headers = ["Category", "Total", "PASS", "PARTIAL", "FAIL", "SKIP", "Pass Rate"]
    for ci, h in enumerate(cat_headers, start=1):
        cell = ws2.cell(row_off, ci, h)
        cell.font  = _font(bold=True, color="FFFFFFFF")
        cell.fill  = _fill(C_MID)
        cell.border = _thin()
        cell.alignment = _align("center", "center", wrap=False)
    row_off += 1

    for cat, counts in summary["by_category"].items():
        cat_total = counts.get("total", 0)
        pass_rate = f"{counts.get('PASS', 0)/cat_total*100:.0f}%" if cat_total else "0%"
        row_vals = [
            cat,
            cat_total,
            counts.get("PASS", 0),
            counts.get("PARTIAL", 0),
            counts.get("FAIL", 0),
            counts.get("SKIP", 0),
            pass_rate,
        ]
        for ci, val in enumerate(row_vals, start=1):
            cell = ws2.cell(row_off, ci, val)
            cell.border = _thin()
            cell.alignment = _align("center", "center", wrap=False)
            cell.fill = _fill(C_ALT if row_off % 2 == 0 else C_WHITE)
        row_off += 1

    for col in ["A", "B", "C", "D", "E", "F", "G"]:
        ws2.column_dimensions[col].width = 22

    wb.save(REPORT_FILE)
    return REPORT_FILE


# ── Main ──────────────────────────────────────────────────────────────────────

async def main():
    anthropic  = get_anthropic()
    cases      = load_test_cases()
    total      = len(cases)

    print(f"\nNEP Analytics — Suite Test Runner")
    print(f"  Test cases : {total}")
    print(f"  API        : {API_BASE}")
    print(f"  Judge model: {JUDGE_MODEL}\n")

    # ── Step 1: Run API calls ─────────────────────────────────────────────────
    print("── Calling API ──────────────────────────────────────────────────────")
    sem = asyncio.Semaphore(MAX_CONCURRENT)
    async with httpx.AsyncClient() as client:
        tasks = [ask_api(client, sem, c) for c in cases]
        api_results = []
        for coro in asyncio.as_completed(tasks):
            res = await coro
            status_icon = "✓" if res["http_status"] == 200 else "✗"
            print(f"  [{status_icon}] {res['id']:4s}  {res['latency_ms']:5d}ms  {res['question'][:60]}")
            api_results.append(res)

    api_results.sort(key=lambda r: r["id"])

    # ── Step 2: LLM Judge ─────────────────────────────────────────────────────
    print("\n── LLM Judge ────────────────────────────────────────────────────────")
    judge_sem = asyncio.Semaphore(MAX_CONCURRENT)
    ctx = (None, judge_sem, anthropic)
    judge_tasks = [judge_one(ctx, r) for r in api_results]
    judged = []
    for coro in asyncio.as_completed(judge_tasks):
        res = await coro
        j = res.get("judge", {})
        icon = {"PASS": "✓", "PARTIAL": "~", "FAIL": "✗"}.get(j.get("verdict", ""), "?")
        print(f"  [{icon}] {res['id']:4s}  {j.get('verdict','?'):8s}  conf={j.get('confidence_pct','?')}%  {res['question'][:50]}")
        judged.append(res)

    judged.sort(key=lambda r: r["id"])

    # ── Step 3: Build summary ─────────────────────────────────────────────────
    overall: dict[str, int] = {"PASS": 0, "PARTIAL": 0, "FAIL": 0, "SKIP": 0}
    by_cat:  dict[str, dict] = {}

    for res in judged:
        verdict  = res.get("judge", {}).get("verdict", "SKIP")
        behavior = res["behavior"]
        cat_label = BEHAVIOR_LABEL.get(behavior, behavior)
        overall[verdict] = overall.get(verdict, 0) + 1
        bc = by_cat.setdefault(cat_label, {"total": 0})
        bc["total"] += 1
        bc[verdict] = bc.get(verdict, 0) + 1

    summary = {"total": total, "overall": overall, "by_category": by_cat}

    # ── Step 4: Write Excel ───────────────────────────────────────────────────
    out = write_excel(judged, summary)

    # ── Print summary ─────────────────────────────────────────────────────────
    print(f"\n{'='*65}")
    print(f"  NEP Suite Test Report — {total} cases")
    print(f"{'='*65}")
    for verdict, cnt in overall.items():
        if cnt:
            pct = cnt / total * 100
            icon = {"PASS": "✓", "PARTIAL": "~", "FAIL": "✗"}.get(verdict, "?")
            print(f"  [{icon}] {verdict:<10}: {cnt:3d}  ({pct:.0f}%)")
    print(f"\n  By category:")
    for cat, counts in by_cat.items():
        p = counts.get("PASS", 0)
        t = counts.get("total", 0)
        print(f"    {cat:<22} {p}/{t} PASS")
    print(f"\n  Report → {out}")
    print(f"{'='*65}\n")


if __name__ == "__main__":
    asyncio.run(main())
