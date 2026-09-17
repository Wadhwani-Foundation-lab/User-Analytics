"""
LLM-as-Judge for NEP Analytics test results.

Evaluates every single-turn answer and every multi-turn conversation turn
using Claude as an independent reviewer. Produces:
  - nep_analytics/tests/judge_report.json   (raw verdicts)
  - nep_analytics/tests/NEP_Judge_Report.xlsx (formatted Excel)

Matches the format of docs/NEP_Judge_Report.xlsx exactly.

Usage:
    python nep_analytics/tests/run_judge.py
"""
from __future__ import annotations

import asyncio
import json
import re
import statistics
import sys
from pathlib import Path

import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

HERE = Path(__file__).resolve().parent
REPO_ROOT = HERE.parent.parent

# ── Add repo root to path so nep_analytics imports work ──────────────────────
sys.path.insert(0, str(REPO_ROOT))

from nep_analytics.core.config import get_anthropic          # noqa: E402
from nep_analytics.core.schema_context import TABLES, SQL_RULES  # noqa: E402

TEST_REPORT = HERE / "test_report.json"
JUDGE_REPORT = HERE / "judge_report.json"
XLSX_OUT = HERE / "NEP_Judge_Report.xlsx"

JUDGE_MODEL = "claude-opus-5"
MAX_CONCURRENT = 4
MAX_ANSWER_LEN = 1500   # truncate long answers in judge prompt
JUDGE_MAX_TOKENS = 4096  # judge writes long reasoning/recommended_sql for complex questions;
                          # 2000 truncated ~25% of single-turn calls mid-JSON, producing
                          # spurious UNCERTAIN verdicts with blank reasoning

# ── Truncation-salvage: recover a real verdict from a cut-off JSON response ──
# The judge schema orders fields verdict -> confidence_pct -> reasoning ->
# sql_issues -> data_accuracy_note -> recommended_sql -> recommended_response.
# When max_tokens cuts the response off, it's almost always the last one or
# two (longest) fields that get truncated — the earlier fields are usually
# intact even though the JSON as a whole no longer parses. Regex-extract them
# directly rather than discarding the whole response.
_SALVAGE_FIELD_RE = {
    "verdict": re.compile(r'"verdict"\s*:\s*"([^"]*)"'),
    "confidence_pct": re.compile(r'"confidence_pct"\s*:\s*(\d+)'),
    "reasoning": re.compile(r'"reasoning"\s*:\s*"((?:[^"\\]|\\.)*)"'),
    "sql_issues": re.compile(r'"sql_issues"\s*:\s*"((?:[^"\\]|\\.)*)"'),
    "data_accuracy_note": re.compile(r'"data_accuracy_note"\s*:\s*"((?:[^"\\]|\\.)*)"'),
    "recommended_sql": re.compile(r'"recommended_sql"\s*:\s*"((?:[^"\\]|\\.)*)"'),
    "recommended_response": re.compile(r'"recommended_response"\s*:\s*"((?:[^"\\]|\\.)*)"'),
}


def _unescape_json_string(s: str) -> str:
    return s.replace('\\n', '\n').replace('\\"', '"').replace('\\\\', '\\')


def _salvage_partial_json(raw: str) -> dict:
    """Best-effort field extraction from a truncated/unparseable judge response."""
    result: dict = {}
    for key, pattern in _SALVAGE_FIELD_RE.items():
        m = pattern.search(raw)
        if not m:
            continue
        result[key] = int(m.group(1)) if key == "confidence_pct" else _unescape_json_string(m.group(1))
    return result

# ── Palette (mirrors generate_report.py) ─────────────────────────────────────
C_HEADER_DARK  = "002C5282"
C_HEADER_MID   = "001E3A5F"
C_HEADER_LIGHT = "00EBF5FF"
C_ROW_ALT      = "00F7F9FC"
C_PASS         = "00D4EDDA"
C_WARN         = "00FFF3CD"
C_FAIL         = "00F8D7DA"
C_INFO         = "00D1ECF1"
C_GREY         = "00E9ECEF"
C_WHITE        = "00FFFFFF"
C_NONE         = "00000000"

VERDICT_FILL = {
    "CORRECT":            C_PASS,
    "PARTIALLY_CORRECT":  C_WARN,
    "INCORRECT":          C_FAIL,
    "UNCERTAIN":          C_INFO,
    "NOT_EVALUABLE":      C_GREY,
}

STATUS_FILL = {
    "PASS":                 C_PASS,
    "DATA_NOT_AVAILABLE":   C_WARN,
    "CLARIFICATION_NEEDED": C_INFO,
    "SQL_ERROR":            C_FAIL,
    "PARSE_ERROR":          C_FAIL,
    "API_ERROR":            C_FAIL,
}


# ── Style helpers ─────────────────────────────────────────────────────────────

def _fill(hex_color: str) -> PatternFill:
    return PatternFill("solid", fgColor=hex_color)

def _font(bold=False, color="FF000000", size=10) -> Font:
    return Font(bold=bold, color=color, size=size, name="Calibri")

def _border() -> Border:
    s = Side(style="thin", color="FFD0D0D0")
    return Border(left=s, right=s, top=s, bottom=s)

def _hdr_font() -> Font:
    return Font(bold=True, color="FFFFFFFF", size=10, name="Calibri")

def _cell(ws, r, c, val, bg=None, bold=False, ha="left", wrap=True):
    cell = ws.cell(row=r, column=c, value=val)
    if bg:
        cell.fill = _fill(bg)
    cell.font = _font(bold=bold)
    cell.alignment = Alignment(horizontal=ha, vertical="top", wrap_text=wrap)
    cell.border = _border()
    return cell


# ── Judge prompt ──────────────────────────────────────────────────────────────

_SCHEMA_SNIPPET = f"""
DATABASE SCHEMA (verified against live DB):
{TABLES}

SQL RULES (must be followed):
{SQL_RULES}
""".strip()

_JUDGE_SYSTEM = f"""You are an expert SQL and analytics reviewer evaluating an AI analytics assistant's responses.
Your job is to assess whether the SQL generated and the answer given are correct for the question asked.

{_SCHEMA_SNIPPET}

Return ONLY a valid JSON object with exactly these keys (no markdown, no extra text):
{{
  "verdict": "CORRECT|PARTIALLY_CORRECT|INCORRECT|UNCERTAIN|NOT_EVALUABLE",
  "confidence_pct": <integer 0-100>,
  "reasoning": "<detailed reasoning — SQL logic, column/table correctness, answer quality>",
  "sql_issues": "<specific SQL problems found, or empty string if none>",
  "data_accuracy_note": "<notes on answer accuracy/reliability, or empty string>",
  "recommended_sql": "<corrected SQL if issues found, or empty string>",
  "recommended_response": "<better response text if needed, or empty string>"
}}

Verdict guidelines:
- CORRECT          SQL logic is sound, columns/tables are right, answer is appropriate.
- PARTIALLY_CORRECT SQL has minor issues (suboptimal but produces right result) or answer is slightly imprecise.
- INCORRECT        SQL has a logical error, wrong column/table, violates a SQL rule, or answer is clearly wrong.
- UNCERTAIN        Cannot verify correctness without running the query against live data; no obvious flaw.
- NOT_EVALUABLE    No SQL was generated (clarification/empty), or the status is SQL_ERROR/API_ERROR with no answer.

When evaluating SQL:
- Check column names against the schema (typos are intentional in the DB — e.g. activity_tittle, particpant_country).
- Check join keys: user_id (user table) ↔ userid (activity table, no underscore).
- Check SQL rules: no CTEs, LIMIT 500, no NOW()/CURRENT_DATE, month_year_order::INTEGER cast, etc.
- Check filter values match exact enum strings from the schema.
- For DATA_NOT_AVAILABLE: the SQL may be correct even with zero rows — evaluate the SQL independently.
"""


def _judge_prompt(question: str, status: str, response_type: str,
                   sql: str, answer: str, row_count: int | None,
                   error: str | None, prior_context: str = "") -> str:
    answer_trunc = (answer or "")[:MAX_ANSWER_LEN]
    if len(answer or "") > MAX_ANSWER_LEN:
        answer_trunc += "... [truncated]"

    parts = []
    if prior_context:
        parts.append(f"PRIOR CONVERSATION CONTEXT:\n{prior_context}\n")
    parts.append(f"QUESTION: {question}")
    parts.append(f"STATUS: {status}")
    parts.append(f"RESPONSE TYPE: {response_type or 'n/a'}")
    parts.append(f"ROW COUNT: {row_count if row_count is not None else 'n/a'}")
    parts.append(f"SQL USED:\n{sql or '(none)'}")
    parts.append(f"ANSWER:\n{answer_trunc or '(none)'}")
    if error:
        parts.append(f"ERROR: {error}")
    parts.append("\nEvaluate the above and return JSON only.")
    return "\n\n".join(parts)


# ── Judge executor ────────────────────────────────────────────────────────────

async def judge_one(semaphore: asyncio.Semaphore, question: str, status: str,
                    response_type: str, sql: str, answer: str,
                    row_count: int | None, error: str | None,
                    prior_context: str = "") -> dict:
    client = get_anthropic()
    user_msg = _judge_prompt(question, status, response_type, sql, answer,
                              row_count, error, prior_context)
    async with semaphore:
        loop = asyncio.get_event_loop()
        resp = await loop.run_in_executor(None, lambda: client.messages.create(
            model=JUDGE_MODEL,
            max_tokens=JUDGE_MAX_TOKENS,
            system=_JUDGE_SYSTEM,
            messages=[{"role": "user", "content": user_msg}],
        ))

    raw = next((b.text for b in resp.content if b.type == "text"), "").strip()
    # Strip markdown fences if present
    if raw.startswith("```"):
        raw = raw.split("```")[1]
        if raw.startswith("json"):
            raw = raw[4:]
    try:
        data = json.loads(raw)
    except json.JSONDecodeError:
        # Fallback: extract first {...} block
        start = raw.find("{")
        end = raw.rfind("}") + 1
        try:
            data = json.loads(raw[start:end]) if start >= 0 and end > start else {}
            if not data:
                raise json.JSONDecodeError("empty extraction", raw, 0)
        except json.JSONDecodeError:
            data = _salvage_partial_json(raw)
            if not data.get("verdict"):
                data["verdict"] = "UNCERTAIN"
                data["reasoning"] = (
                    f"Judge response was truncated (stop_reason={resp.stop_reason}) and "
                    f"unparseable even after salvage (first 200 chars): {raw[:200]}"
                )

    return {
        "verdict":             data.get("verdict", "UNCERTAIN"),
        "confidence_pct":      data.get("confidence_pct", 50),
        "reasoning":           data.get("reasoning", ""),
        "sql_issues":          data.get("sql_issues", ""),
        "data_accuracy_note":  data.get("data_accuracy_note", ""),
        "recommended_sql":     data.get("recommended_sql", ""),
        "recommended_response":data.get("recommended_response", ""),
    }


# ── Single-turn judging ───────────────────────────────────────────────────────

async def judge_single_turn(results: list[dict]) -> list[dict]:
    semaphore = asyncio.Semaphore(MAX_CONCURRENT)
    judged = []

    async def _judge(r: dict) -> dict:
        print(f"  Judging Q{r['num']:02d} — {r['question'][:60]}")
        j = await judge_one(
            semaphore,
            question=r["question"],
            status=r["status"],
            response_type=r.get("response_type", ""),
            sql=r.get("sql_used", ""),
            answer=r.get("answer", ""),
            row_count=r.get("row_count"),
            error=r.get("error"),
        )
        return {**r, "judge": j}

    tasks = [_judge(r) for r in results]
    judged = await asyncio.gather(*tasks)
    return list(judged)


# ── Multi-turn judging ────────────────────────────────────────────────────────

async def judge_multi_turn(scenario_results: list[dict]) -> list[dict]:
    semaphore = asyncio.Semaphore(MAX_CONCURRENT)
    judged_scenarios = []

    for scenario in scenario_results:
        print(f"\n  Scenario {scenario['id']}: {scenario['name']}")
        judged_turns = []
        prior_ctx_parts = []

        async def _judge_turn(turn: dict, prior: str) -> dict:
            print(f"    T{turn['turn_num']} — {turn['question'][:55]}")
            j = await judge_one(
                semaphore,
                question=turn["question"],
                status=turn["status"],
                response_type=turn.get("response_type", ""),
                sql=turn.get("sql_used", ""),
                answer=turn.get("answer", ""),
                row_count=turn.get("row_count"),
                error=turn.get("error"),
                prior_context=prior,
            )
            return {**turn, "judge": j}

        # Run turns sequentially within a scenario (to preserve context order)
        for turn in scenario["turns"]:
            prior_ctx = "\n".join(prior_ctx_parts) if prior_ctx_parts else ""
            judged_turn = await _judge_turn(turn, prior_ctx)
            judged_turns.append(judged_turn)
            prior_ctx_parts.append(
                f"Q: {turn['question']}\nA: {(turn.get('answer') or '')[:400]}"
            )

        judged_scenarios.append({**scenario, "turns": judged_turns})

    return judged_scenarios


# ── Excel generation ──────────────────────────────────────────────────────────

ALL_Q_COLS = [
    "#", "Category", "Question", "Status", "Path",
    "Certified Metric", "Confidence", "Response Type",
    "Answer", "SQL Used", "Rows", "Latency (ms)", "Error",
    None, None,   # spacer cols 14-15
    "Judge Verdict", "Judge Confidence %", "Reasoning",
    "SQL Issues", "Data Accuracy Note", "Recommended SQL", "Recommended Response",
]
ALL_Q_WIDTHS = [5, 28, 55, 12, 12, 22, 11, 14, 70, 58, 6, 12, 40, 2, 2,
                18, 14, 80, 55, 55, 70, 70]

MT_COLS = [
    "Scenario", "Turn #", "Question", "Status", "Response Type",
    "Answer", "SQL Used", "Rows", "Latency (ms)",
    None,   # spacer col 10
    "Judge Verdict", "Judge Confidence %", "Reasoning",
    "SQL Issues", "Data Accuracy Note",
]
MT_WIDTHS = [10, 7, 55, 12, 14, 70, 58, 6, 12, 2, 18, 14, 80, 55, 55]


def _section_hdr(ws, row, title, ncols):
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=ncols)
    c = ws.cell(row=row, column=1, value=title)
    c.fill = _fill(C_HEADER_DARK)
    c.font = Font(bold=True, color="FFFFFFFF", size=11, name="Calibri")
    c.alignment = Alignment(horizontal="left", vertical="center")
    ws.row_dimensions[row].height = 20


def _col_hdr_row(ws, row, cols):
    for ci, h in enumerate(cols, 1):
        c = ws.cell(row=row, column=ci, value=h or "")
        c.fill = _fill(C_HEADER_MID if h else C_WHITE)
        c.font = _hdr_font()
        c.alignment = Alignment(horizontal="center", vertical="center")
        c.border = _border()
    ws.row_dimensions[row].height = 16


def _apply_widths(ws, widths):
    for ci, w in enumerate(widths, 1):
        ws.column_dimensions[openpyxl.utils.get_column_letter(ci)].width = w


def build_all_questions_sheet(wb, judged_results: list[dict], run_at: str):
    ws = wb.create_sheet("All Questions")
    ncols = len(ALL_Q_COLS)
    _apply_widths(ws, ALL_Q_WIDTHS)

    title = (f"NEP Analytics — LLM-as-Judge Report   |   Run: {run_at}   |   "
             f"{len(judged_results)} questions   |   Judge: {JUDGE_MODEL}")
    _section_hdr(ws, 1, title, ncols)
    _col_hdr_row(ws, 2, ALL_Q_COLS)
    ws.freeze_panes = "A3"
    ws.auto_filter.ref = f"A2:{openpyxl.utils.get_column_letter(ncols)}2"

    for ri, r in enumerate(judged_results, 3):
        status = r["status"]
        j = r.get("judge", {})
        verdict = j.get("verdict", "UNCERTAIN")
        row_bg = C_WHITE if ri % 2 == 1 else "00F2F6FA"

        values = [
            r["num"],
            r["category"],
            r["question"],
            status,
            "certified" if r.get("certified") else "sql_gen",
            r.get("metric") or "",
            "",
            r.get("response_type") or "",
            r.get("answer") or "",
            r.get("sql_used") or "",
            r.get("row_count") or "",
            r.get("latency_ms") or "",
            r.get("error") or "",
            "", "",   # spacers
            verdict,
            j.get("confidence_pct", ""),
            j.get("reasoning", ""),
            j.get("sql_issues", ""),
            j.get("data_accuracy_note", ""),
            j.get("recommended_sql", ""),
            j.get("recommended_response", ""),
        ]

        center_cols = {1, 4, 5, 7, 8, 11, 12, 16, 17}
        for ci, val in enumerate(values, 1):
            col_idx = ci
            if col_idx in (14, 15):
                c = ws.cell(row=ri, column=col_idx, value="")
                c.fill = _fill(C_WHITE)
                continue
            bg = (STATUS_FILL.get(status, C_WARN) if col_idx == 4
                  else VERDICT_FILL.get(verdict, C_GREY) if col_idx == 16
                  else row_bg)
            ha = "center" if col_idx in center_cols else "left"
            _cell(ws, ri, col_idx, val, bg=bg, ha=ha)

        ans_len = len(str(r.get("answer") or ""))
        reason_len = len(str(j.get("reasoning") or ""))
        ws.row_dimensions[ri].height = min(max(18, max(ans_len, reason_len) // 8), 150)


def build_multi_turn_sheet(wb, judged_scenarios: list[dict], run_at: str):
    ws = wb.create_sheet("Multi-Turn Scenarios")
    ncols = len(MT_COLS)
    _apply_widths(ws, MT_WIDTHS)

    total_turns = sum(len(s["turns"]) for s in judged_scenarios)
    title = (f"Multi-Turn Scenarios — LLM-as-Judge   |   Run: {run_at}   |   "
             f"{len(judged_scenarios)} scenarios  {total_turns} turns   |   Judge: {JUDGE_MODEL}")
    _section_hdr(ws, 1, title, ncols)
    _col_hdr_row(ws, 2, MT_COLS)
    ws.freeze_panes = "A3"

    ri = 3
    for scenario in judged_scenarios:
        overall = scenario["overall_status"]
        sc_fill_color = {
            "PASS": C_PASS, "PARTIAL": C_WARN, "ERROR": C_FAIL
        }.get(overall, C_GREY)

        # Scenario header row
        sc_label = f"{scenario['id']}  —  {scenario['name']}"
        for ci in range(1, ncols + 1):
            c = ws.cell(row=ri, column=ci)
            c.fill = _fill("001E3A5F")
            c.font = Font(bold=True, color="FFFFFFFF", size=10, name="Calibri")
            c.alignment = Alignment(horizontal="left", vertical="center")
        ws.cell(row=ri, column=1).value = sc_label
        badge = ws.cell(row=ri, column=4)
        badge.value = overall
        badge.fill = _fill(sc_fill_color)
        badge.font = Font(bold=True, size=10, name="Calibri")
        badge.alignment = Alignment(horizontal="center", vertical="center")
        ws.row_dimensions[ri].height = 16
        ri += 1

        for turn in scenario["turns"]:
            status = turn["status"]
            j = turn.get("judge", {})
            verdict = j.get("verdict", "UNCERTAIN")
            row_bg = "00F7F9FC"

            values = [
                f"{scenario['id']}",
                f"T{turn['turn_num']}",
                turn["question"],
                status,
                turn.get("response_type") or "",
                turn.get("answer") or "",
                turn.get("sql_used") or "",
                turn.get("row_count") or "",
                turn.get("latency_ms") or "",
                "",   # spacer
                verdict,
                j.get("confidence_pct", ""),
                j.get("reasoning", ""),
                j.get("sql_issues", ""),
                j.get("data_accuracy_note", ""),
            ]

            center_cols = {1, 2, 4, 5, 8, 9, 11, 12}
            for ci, val in enumerate(values, 1):
                if ci == 10:
                    ws.cell(row=ri, column=ci, value="").fill = _fill(C_WHITE)
                    continue
                bg = (STATUS_FILL.get(status, C_WARN) if ci == 4
                      else VERDICT_FILL.get(verdict, C_GREY) if ci == 11
                      else row_bg)
                ha = "center" if ci in center_cols else "left"
                _cell(ws, ri, ci, val, bg=bg, ha=ha)

            reason_len = len(str(j.get("reasoning") or ""))
            ws.row_dimensions[ri].height = min(max(18, reason_len // 8), 150)
            ri += 1

        ws.row_dimensions[ri].height = 6
        ri += 1


def build_summary_sheet(wb, judged_results: list[dict],
                         judged_scenarios: list[dict], report: dict):
    ws = wb.create_sheet("Summary")
    ws.column_dimensions["A"].width = 52
    ws.column_dimensions["B"].width = 22
    ws.column_dimensions["C"].width = 18
    ws.column_dimensions["D"].width = 18

    s = report["summary"]
    total = report["meta"]["total_questions"]
    run_at = report["meta"]["finished_at"][:19].replace("T", " ")
    latencies = [r["latency_ms"] for r in judged_results if r.get("latency_ms")]
    avg_lat = round(statistics.mean(latencies) / 1000, 1) if latencies else 0
    p95_lat = round(sorted(latencies)[int(len(latencies) * 0.95)] / 1000, 1) if latencies else 0

    # Verdict tallies
    verdicts = [r.get("judge", {}).get("verdict", "UNCERTAIN") for r in judged_results]
    v_correct   = verdicts.count("CORRECT")
    v_partial   = verdicts.count("PARTIALLY_CORRECT")
    v_incorrect = verdicts.count("INCORRECT")
    v_uncertain = verdicts.count("UNCERTAIN")
    v_na        = verdicts.count("NOT_EVALUABLE")
    conf_vals   = [r.get("judge", {}).get("confidence_pct") for r in judged_results
                   if isinstance(r.get("judge", {}).get("confidence_pct"), (int, float))]
    avg_conf = round(statistics.mean(conf_vals), 1) if conf_vals else 0

    def section_hdr(row, title):
        ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=4)
        c = ws.cell(row=row, column=1, value=title)
        c.fill = _fill(C_HEADER_DARK)
        c.font = Font(bold=True, color="FFFFFFFF", size=11, name="Calibri")
        c.alignment = Alignment(horizontal="left", vertical="center")
        ws.row_dimensions[row].height = 18

    def kv(row, label, value, val_bg=C_NONE):
        lc = ws.cell(row=row, column=1, value=label)
        lc.fill = _fill(C_HEADER_LIGHT)
        lc.font = Font(bold=True, size=10, name="Calibri")
        lc.alignment = Alignment(horizontal="left", vertical="center")
        lc.border = _border()
        vc = ws.cell(row=row, column=2, value=value)
        if val_bg and val_bg != C_NONE:
            vc.fill = _fill(val_bg)
        vc.font = Font(size=10, name="Calibri")
        vc.alignment = Alignment(horizontal="left", vertical="center")
        vc.border = _border()
        ws.row_dimensions[row].height = 16

    section_hdr(1, f"Single-Turn Test Summary   |   Run: {run_at}")
    kv(2,  "Total Questions Tested", total)
    kv(3,  "PASS (system status)",   f"{s['pass']}  ({s['pass_rate_pct']}%)", C_PASS)
    kv(4,  "Data Not Available",     f"{s['data_not_available']}  ({round(s['data_not_available']/total*100,1)}%)", C_WARN)
    kv(5,  "SQL Error",              f"{s['sql_error']}  ({round(s['sql_error']/total*100,1)}%)", C_FAIL if s['sql_error'] else C_NONE)
    kv(6,  "Certified path",         f"{s['certified_path']}  ({round(s['certified_path']/total*100,1)}%)")
    kv(7,  "SQL-gen path",           f"{s['sql_gen_path']}  ({round(s['sql_gen_path']/total*100,1)}%)")
    kv(8,  "Avg response latency",   f"{avg_lat}s")
    kv(9,  "P95 response latency",   f"{p95_lat}s")

    section_hdr(11, f"Judge Verdicts   |   Model: {JUDGE_MODEL}")
    kv(12, "CORRECT",            f"{v_correct}  ({round(v_correct/total*100,1)}%)", C_PASS)
    kv(13, "PARTIALLY_CORRECT",  f"{v_partial}  ({round(v_partial/total*100,1)}%)", C_WARN)
    kv(14, "INCORRECT",          f"{v_incorrect}  ({round(v_incorrect/total*100,1)}%)", C_FAIL if v_incorrect else C_NONE)
    kv(15, "UNCERTAIN",          f"{v_uncertain}  ({round(v_uncertain/total*100,1)}%)", C_INFO)
    kv(16, "NOT_EVALUABLE",      f"{v_na}  ({round(v_na/total*100,1)}%)", C_GREY)
    kv(17, "Avg judge confidence",f"{avg_conf}%")

    # By-category table
    section_hdr(19, "Results by Category")
    for ci, h in enumerate(["Category", "Total", "PASS", "Judge: Correct/Partial/Incorrect"], 1):
        c = ws.cell(row=20, column=ci, value=h)
        c.fill = _fill(C_HEADER_MID)
        c.font = _hdr_font()
        c.alignment = Alignment(horizontal="center", vertical="center")
        c.border = _border()
    ws.row_dimensions[20].height = 16

    by_cat = report["by_category"]
    # Build per-category verdict counts
    cat_verdicts: dict[str, dict] = {}
    for r in judged_results:
        cat = r["category"]
        v = r.get("judge", {}).get("verdict", "UNCERTAIN")
        if cat not in cat_verdicts:
            cat_verdicts[cat] = {"CORRECT": 0, "PARTIALLY_CORRECT": 0, "INCORRECT": 0, "UNCERTAIN": 0, "NOT_EVALUABLE": 0}
        cat_verdicts[cat][v] = cat_verdicts[cat].get(v, 0) + 1

    for ri, (cat, counts) in enumerate(sorted(by_cat.items()), 21):
        pct = round(counts["pass"] / counts["total"] * 100) if counts["total"] else 0
        cv = cat_verdicts.get(cat, {})
        judge_summary = (f"✓{cv.get('CORRECT',0)}  "
                         f"~{cv.get('PARTIALLY_CORRECT',0)}  "
                         f"✗{cv.get('INCORRECT',0)}  "
                         f"?{cv.get('UNCERTAIN',0)}")
        row_bg = C_ROW_ALT if ri % 2 == 0 else C_WHITE
        for ci, val in enumerate([cat, counts["total"], f"{counts['pass']} ({pct}%)", judge_summary], 1):
            c = ws.cell(row=ri, column=ci, value=val)
            c.fill = _fill(row_bg)
            c.font = Font(size=10, name="Calibri")
            c.alignment = Alignment(horizontal="left" if ci in (1, 4) else "center", vertical="center")
            c.border = _border()
        ws.row_dimensions[ri].height = 15

    # Multi-turn summary
    mt_base = 21 + len(by_cat) + 2
    mt = report.get("multi_turn_summary", {})
    if mt and judged_scenarios:
        section_hdr(mt_base, "Multi-Turn Scenarios — Judge Summary")
        mt_verdicts = [t.get("judge", {}).get("verdict", "UNCERTAIN")
                       for s in judged_scenarios for t in s["turns"]]
        mt_total_turns = len(mt_verdicts)
        mt_c = mt_verdicts.count("CORRECT")
        mt_p = mt_verdicts.count("PARTIALLY_CORRECT")
        mt_i = mt_verdicts.count("INCORRECT")
        mt_u = mt_verdicts.count("UNCERTAIN")
        kv(mt_base+1, "Total Scenarios", len(judged_scenarios))
        kv(mt_base+2, "Total Turns Judged", mt_total_turns)
        kv(mt_base+3, "CORRECT turns",           f"{mt_c}  ({round(mt_c/mt_total_turns*100,1) if mt_total_turns else 0}%)", C_PASS)
        kv(mt_base+4, "PARTIALLY_CORRECT turns",  f"{mt_p}  ({round(mt_p/mt_total_turns*100,1) if mt_total_turns else 0}%)", C_WARN)
        kv(mt_base+5, "INCORRECT turns",          f"{mt_i}  ({round(mt_i/mt_total_turns*100,1) if mt_total_turns else 0}%)", C_FAIL if mt_i else C_NONE)
        kv(mt_base+6, "UNCERTAIN turns",          f"{mt_u}  ({round(mt_u/mt_total_turns*100,1) if mt_total_turns else 0}%)", C_INFO)

    # Errors / Issues log
    err_base = mt_base + 10 if mt else mt_base + 2
    section_hdr(err_base, "Questions Flagged by Judge (INCORRECT or PARTIALLY_CORRECT)")
    for ci, h in enumerate(["#", "Question", "Verdict", "SQL Issues"], 1):
        c = ws.cell(row=err_base+1, column=ci, value=h)
        c.fill = _fill(C_HEADER_MID)
        c.font = _hdr_font()
        c.border = _border()
    ws.row_dimensions[err_base+1].height = 15

    flagged = [r for r in judged_results
               if r.get("judge", {}).get("verdict") in ("INCORRECT", "PARTIALLY_CORRECT")]
    if flagged:
        for ri, r in enumerate(flagged, err_base+2):
            j = r.get("judge", {})
            v = j.get("verdict", "")
            for ci, val in enumerate([r["num"], r["question"], v, j.get("sql_issues","")[:200]], 1):
                c = ws.cell(row=ri, column=ci, value=val)
                c.fill = _fill(VERDICT_FILL.get(v, C_GREY))
                c.font = Font(size=10, name="Calibri")
                c.alignment = Alignment(horizontal="left" if ci != 1 else "center",
                                        vertical="top", wrap_text=True)
                c.border = _border()
            ws.row_dimensions[ri].height = 30
    else:
        c = ws.cell(row=err_base+2, column=1, value="No INCORRECT / PARTIALLY_CORRECT verdicts")
        c.fill = _fill(C_PASS)
        c.font = Font(bold=True, size=10, name="Calibri", color="FF155724")
        c.border = _border()


# ── Main ──────────────────────────────────────────────────────────────────────

async def main_async():
    with open(TEST_REPORT, encoding="utf-8") as f:
        report = json.load(f)

    results = report["results"]
    scenario_results = report.get("scenario_results", [])
    run_at = report["meta"]["finished_at"][:16].replace("T", " ")

    print(f"\nNEP Analytics — LLM-as-Judge")
    print(f"  Judge model  : {JUDGE_MODEL}")
    print(f"  Single-turn  : {len(results)} questions")
    print(f"  Multi-turn   : {len(scenario_results)} scenarios  "
          f"({sum(len(s['turns']) for s in scenario_results)} turns)")
    print(f"  Concurrency  : {MAX_CONCURRENT}")
    print()

    print("── Single-turn judgement ──────────────────────────────────────────")
    judged_results = await judge_single_turn(results)

    print("\n── Multi-turn judgement ───────────────────────────────────────────")
    judged_scenarios = await judge_multi_turn(scenario_results)

    # Save raw judge report
    judge_data = {
        "meta": report["meta"],
        "judge_model": JUDGE_MODEL,
        "summary": report["summary"],
        "multi_turn_summary": report.get("multi_turn_summary", {}),
        "by_category": report["by_category"],
        "results": judged_results,
        "scenario_results": judged_scenarios,
    }
    with open(JUDGE_REPORT, "w", encoding="utf-8") as f:
        json.dump(judge_data, f, indent=2, ensure_ascii=False)
    print(f"\nRaw judge data saved → {JUDGE_REPORT}")

    # Generate Excel
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    build_all_questions_sheet(wb, judged_results, run_at)
    build_multi_turn_sheet(wb, judged_scenarios, run_at)
    build_summary_sheet(wb, judged_results, judged_scenarios, judge_data)
    wb.save(XLSX_OUT)

    # Print verdict summary
    verdicts = [r.get("judge", {}).get("verdict", "?") for r in judged_results]
    total = len(verdicts)
    print(f"\n{'='*65}")
    print(f"  Judge Verdicts — {total} single-turn questions")
    print(f"{'='*65}")
    for v, label in [("CORRECT","CORRECT"), ("PARTIALLY_CORRECT","PARTIALLY_CORRECT"),
                      ("INCORRECT","INCORRECT"), ("UNCERTAIN","UNCERTAIN"),
                      ("NOT_EVALUABLE","NOT_EVALUABLE")]:
        n = verdicts.count(v)
        print(f"  {label:<22}: {n:>3}  ({round(n/total*100,1)}%)")
    mt_turns = [t.get("judge", {}).get("verdict", "?")
                for s in judged_scenarios for t in s["turns"]]
    if mt_turns:
        print(f"\n  Multi-turn turns judged: {len(mt_turns)}")
        for v in ("CORRECT", "PARTIALLY_CORRECT", "INCORRECT", "UNCERTAIN"):
            n = mt_turns.count(v)
            print(f"    {v:<22}: {n}  ({round(n/len(mt_turns)*100,1)}%)")
    print(f"\n  Report saved → {XLSX_OUT}")
    print(f"{'='*65}")


def main():
    asyncio.run(main_async())


if __name__ == "__main__":
    main()
