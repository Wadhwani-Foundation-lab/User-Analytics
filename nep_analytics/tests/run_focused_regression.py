"""
Focused Regression Test Runner
Runs only the questions relevant to the 4 targeted schema_context.py fixes
plus a small set of regression guards, then produces a single consolidated
XLSX report.

Fix targets:
  Fix 1 — Rule 30 (B3):  "How many users have had AI conversations?"
  Fix 2 — TABLES (E1):   event_id in liftoffx → use events table instead
  Fix 3 — Rule 23 (D5):  scalar "last N months" must use literal date range
  Fix 4 — Rule 22 (C1, Q39, MT10 T2): comparison queries never filter to one value

Regression guards (previously passing — must stay passing):
  Q02  week-over-week trend  (Rule 23: CASE A — still uses precomputed columns)
  MT11 last-6-months scalar  (Rule 23: CASE B — but user table uses created_datetime)
  A1, A2, B4 from suite

Usage (from repo root):
    python nep_analytics/tests/run_focused_regression.py
"""
from __future__ import annotations

import asyncio
import json
import os
import sys
import time
import uuid
from pathlib import Path

import anthropic
import httpx
import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

HERE = Path(__file__).resolve().parent
REPO_ROOT = HERE.parent.parent
sys.path.insert(0, str(REPO_ROOT))

API_BASE = os.getenv("API_BASE_URL", "http://localhost:8002")
API_KEY  = os.getenv("API_KEY", "nep-analytics-internal-2026")
HEADERS  = {"Content-Type": "application/json", "x-api-key": API_KEY}
TIMEOUT  = 90.0

JUDGE_MODEL = "claude-opus-5"

REPORT_PATH = HERE / "NEP_Focused_Regression_Report.xlsx"

SUITE_FILE = HERE / "nep_query_test_suite.xlsx"
SCENARIOS_FILE = HERE / "multi_turn_scenarios.json"
QUESTIONS_FILE = REPO_ROOT / "docs" / "NEP Test Questions.xlsx"


# ── Test plan ─────────────────────────────────────────────────────────────────

# Single-turn questions from the 78-question set (question numbers to run)
SINGLE_TURN_IDS = {
    39: "Fix 4 target — Rule 22 comparison",
    2:  "Regression guard — Rule 23 CASE A (trend, precomputed columns)",
}

# Multi-turn scenarios to run (id → note)
MT_IDS = {
    "MT10": "Fix 4 target — Rule 22 comparison, T2",
    "MT11": "Regression guard — Rule 23 CASE B + Rule 31 engagement",
}

# Suite question IDs to run (id → note)
SUITE_IDS = {
    "B3": "Fix 1 target — Rule 30 scalar 'AI conversations'",
    "C1": "Fix 4 target — Rule 22 ATTENDED vs NOSHOW comparison",
    "D5": "Fix 3 target — Rule 23 scalar 'last 3 months'",
    "E1": "Fix 2 target — event_id in liftoffx",
    "A1": "Regression guard — active mentors PUBLIC (scalar)",
    "A2": "Regression guard — expert mentors by program (breakdown)",
    "B4": "Regression guard — signup counts by month (series)",
}


# ── Palette ───────────────────────────────────────────────────────────────────

C_DARK  = "002C5282"
C_MID   = "001E3A5F"
C_LIGHT = "00EBF5FF"
C_ALT   = "00F7F9FC"
C_WHITE = "00FFFFFF"
C_OK    = "00D4EDDA"
C_PASS  = "00D4EDDA"
C_WARN  = "00FFF3CD"
C_ERR   = "00F8D7DA"
C_PART  = "00FFE8A1"
C_UNCRT = "00F0E6FF"

VERDICT_FILL = {
    "CORRECT":          C_OK,
    "PASS":             C_PASS,
    "PARTIALLY_CORRECT": C_PART,
    "UNCERTAIN":        C_UNCRT,
    "INCORRECT":        C_ERR,
    "FAIL":             C_ERR,
    "DATA_NOT_AVAILABLE": C_WARN,
    "SQL_ERROR":        C_ERR,
    "API_ERROR":        C_ERR,
    "CLARIFICATION_NEEDED": C_LIGHT,
    "ERROR":            C_ERR,
}


# ── Loaders ───────────────────────────────────────────────────────────────────

def load_single_turn_questions() -> list[dict]:
    wb = openpyxl.load_workbook(QUESTIONS_FILE)
    ws = wb.active
    out = []
    for row in ws.iter_rows(min_row=2, values_only=True):
        num = row[0]
        if num in SINGLE_TURN_IDS:
            out.append({
                "num":      num,
                "category": str(row[1] or "").strip(),
                "question": str(row[2]).strip(),
                "note":     SINGLE_TURN_IDS[num],
            })
    return out


def load_mt_scenarios() -> list[dict]:
    with open(SCENARIOS_FILE, encoding="utf-8") as f:
        all_scens = json.load(f)
    return [s for s in all_scens if s["id"] in MT_IDS]


def load_suite_questions() -> list[dict]:
    wb = openpyxl.load_workbook(SUITE_FILE)
    ws = wb["Test Cases"]
    out = []
    for r in range(5, ws.max_row + 1):
        row = [ws.cell(r, c).value for c in range(1, ws.max_column + 1)]
        if not row[0]:
            continue
        cid = str(row[0]).strip()
        if cid in SUITE_IDS:
            out.append({
                "id":       cid,
                "category": str(row[1] or "").strip(),
                "question": str(row[2] or "").strip(),
                "note":     SUITE_IDS[cid],
            })
    return out


# ── API helpers ───────────────────────────────────────────────────────────────

def _classify(answer: str, sql: str, response_type: str, table: dict, chart: dict,
              error: str | None, http_status: int) -> str:
    if error or http_status >= 500:
        return "SQL_ERROR" if (error and "sql" in error.lower()) else "API_ERROR"
    if not sql:
        return "CLARIFICATION_NEEDED"
    ans = (answer or "").lower()
    no_data = ["no data", "no results", "no matching", "no records", "0 results",
               "zero results", "not find", "no users found", "no events found",
               "no mentors", "does not exist", "no information available", "no entries"]
    if any(p in ans for p in no_data):
        return "DATA_NOT_AVAILABLE"
    if response_type == "table" and isinstance(table.get("rows"), list) and len(table["rows"]) == 0:
        return "DATA_NOT_AVAILABLE"
    if response_type in ("bar_chart","line_chart") and isinstance(chart.get("labels"), list) and len(chart["labels"]) == 0:
        return "DATA_NOT_AVAILABLE"
    return "PASS"


async def call_api(client: httpx.AsyncClient, sem: asyncio.Semaphore,
                   question: str, session_id: str, history: list) -> dict:
    payload = {"session_id": session_id, "question": question, "history": history}
    t0 = time.perf_counter()
    error = None
    http_status = 0
    resp_data: dict = {}

    try:
        async with sem:
            r = await client.post(f"{API_BASE}/api/chat", json=payload,
                                  headers=HEADERS, timeout=TIMEOUT)
            http_status = r.status_code
            if http_status == 200:
                resp_data = r.json()
            else:
                error = f"HTTP {http_status}: {r.text[:200]}"
    except Exception as exc:
        error = str(exc)[:200]

    latency_ms = int((time.perf_counter() - t0) * 1000)
    answer   = resp_data.get("answer", "")
    sql      = resp_data.get("sql_used", "") or ""
    rtype    = resp_data.get("response_type", "text")
    table    = resp_data.get("table_data") or {}
    chart    = resp_data.get("chart_config") or {}
    row_cnt  = (len(table["rows"]) if isinstance(table.get("rows"), list)
                else len(chart["labels"]) if isinstance(chart.get("labels"), list)
                else None)

    status = _classify(answer, sql, rtype, table, chart, error, http_status)
    return {
        "latency_ms": latency_ms,
        "status":     status,
        "answer":     answer,
        "sql":        sql,
        "rtype":      rtype,
        "row_count":  row_cnt,
        "error":      error,
    }


# ── LLM Judge ─────────────────────────────────────────────────────────────────

_JUDGE_SYSTEM = """\
You are an expert SQL and analytics reviewer. Given:
- A natural-language question
- The SQL that was generated
- The answer returned

Return ONLY a JSON object with exactly these keys:
{
  "verdict": "CORRECT" | "PARTIALLY_CORRECT" | "INCORRECT" | "UNCERTAIN",
  "confidence": <0-100 integer>,
  "reason": "<one concise sentence explaining the verdict>"
}

Scoring guide:
- CORRECT:           SQL targets the right tables/columns; answer addresses the question well.
- PARTIALLY_CORRECT: SQL is directionally right but has a minor flaw (wrong filter, missing column).
- INCORRECT:         SQL has a fundamental error (wrong table, missing join, wrong GROUP BY).
- UNCERTAIN:         Not enough information to judge (e.g. result is 0 and may be legitimately empty).

Focus on SQL correctness, not prose quality. Return ONLY the JSON object — no markdown.
"""


def judge(question: str, sql: str, answer: str, status: str) -> dict:
    if status in ("API_ERROR", "SQL_ERROR", "CLARIFICATION_NEEDED"):
        return {"verdict": status, "confidence": 100,
                "reason": f"Test returned {status} — no SQL to evaluate."}
    if status == "DATA_NOT_AVAILABLE" and not sql:
        return {"verdict": "UNCERTAIN", "confidence": 50,
                "reason": "No SQL was generated — cannot evaluate."}

    client = anthropic.Anthropic()
    user_msg = (
        f"Question: {question}\n\n"
        f"SQL:\n{sql or '(none)'}\n\n"
        f"Answer: {answer or '(none)'}\n\n"
        f"API status: {status}"
    )
    try:
        resp = client.messages.create(
            model=JUDGE_MODEL,
            max_tokens=1000,
            system=_JUDGE_SYSTEM,
            messages=[{"role": "user", "content": user_msg}],
        )
        raw = next((b.text for b in resp.content if b.type == "text"), "").strip()
        # strip markdown fence if present
        if raw.startswith("```"):
            raw = "\n".join(raw.split("\n")[1:])
            if raw.endswith("```"):
                raw = raw[:-3].strip()
        return json.loads(raw)
    except Exception as exc:
        return {"verdict": "UNCERTAIN", "confidence": 0, "reason": f"Judge error: {exc}"}


# ── Run single-turn 78Q tests ─────────────────────────────────────────────────

async def run_single_turn(questions: list[dict]) -> list[dict]:
    sem = asyncio.Semaphore(3)
    results = []

    async def _run(q: dict):
        session_id = f"regr-78q-{q['num']}-{uuid.uuid4().hex[:6]}"
        r = await call_api(httpx.AsyncClient(), sem, q["question"], session_id, [])
        verdict = judge(q["question"], r["sql"], r["answer"], r["status"])
        result = {
            "source": "78-question set",
            "id":     f"Q{q['num']:02d}",
            "note":   q["note"],
            "question": q["question"],
            "api_status": r["status"],
            "latency_ms": r["latency_ms"],
            "row_count":  r["row_count"],
            "verdict":    verdict.get("verdict", "UNCERTAIN"),
            "confidence": verdict.get("confidence", 0),
            "reason":     verdict.get("reason", ""),
            "sql":        r["sql"],
            "answer":     r["answer"],
        }
        icon = {"CORRECT": "✓", "PARTIALLY_CORRECT": "~",
                "INCORRECT": "✗", "UNCERTAIN": "?"}.get(result["verdict"], " ")
        print(f"  [{icon}] {result['id']:6s}  {result['verdict']:<20s}  {q['note'][:55]}")
        results.append(result)

    async with httpx.AsyncClient() as client_unused:
        tasks = [_run(q) for q in questions]
        await asyncio.gather(*tasks)

    return results


# ── Run multi-turn scenarios ──────────────────────────────────────────────────

async def run_multi_turn(scenarios: list[dict]) -> list[dict]:
    sem = asyncio.Semaphore(2)
    results = []

    async def _run_scenario(sc: dict):
        session_id = f"regr-mt-{sc['id']}-{uuid.uuid4().hex[:6]}"
        history: list[dict] = []
        sc_note = MT_IDS.get(sc["id"], "")
        sc_results = []

        async with sem:
            for ti, turn_q in enumerate(sc["turns"], 1):
                async with httpx.AsyncClient() as client:
                    r = await call_api(client, asyncio.Semaphore(1), turn_q, session_id, history)

                verdict = judge(turn_q, r["sql"], r["answer"], r["status"])
                row = {
                    "source": "Multi-turn",
                    "id":     f"{sc['id']}-T{ti}",
                    "note":   sc_note if ti == 1 else "",
                    "question": turn_q,
                    "api_status": r["status"],
                    "latency_ms": r["latency_ms"],
                    "row_count":  r["row_count"],
                    "verdict":    verdict.get("verdict", "UNCERTAIN"),
                    "confidence": verdict.get("confidence", 0),
                    "reason":     verdict.get("reason", ""),
                    "sql":        r["sql"],
                    "answer":     r["answer"],
                }
                sc_results.append(row)
                icon = {"CORRECT": "✓", "PARTIALLY_CORRECT": "~",
                        "INCORRECT": "✗", "UNCERTAIN": "?"}.get(row["verdict"], " ")
                print(f"  [{icon}] {row['id']:8s}  {row['verdict']:<20s}  {turn_q[:55]}")

                # Build history for next turn
                history.append({"role": "user",    "content": turn_q})
                history.append({"role": "assistant","content": r["answer"]})

        results.extend(sc_results)

    await asyncio.gather(*[_run_scenario(sc) for sc in scenarios])
    return results


# ── Run suite questions ───────────────────────────────────────────────────────

async def run_suite(questions: list[dict]) -> list[dict]:
    sem = asyncio.Semaphore(3)
    results = []

    async def _run(q: dict):
        session_id = f"regr-suite-{q['id']}-{uuid.uuid4().hex[:6]}"
        async with httpx.AsyncClient() as client:
            r = await call_api(client, sem, q["question"], session_id, [])
        verdict = judge(q["question"], r["sql"], r["answer"], r["status"])
        result = {
            "source":     "24-question suite",
            "id":         q["id"],
            "note":       q["note"],
            "question":   q["question"],
            "api_status": r["status"],
            "latency_ms": r["latency_ms"],
            "row_count":  r["row_count"],
            "verdict":    verdict.get("verdict", "UNCERTAIN"),
            "confidence": verdict.get("confidence", 0),
            "reason":     verdict.get("reason", ""),
            "sql":        r["sql"],
            "answer":     r["answer"],
        }
        icon = {"CORRECT": "✓", "PARTIALLY_CORRECT": "~",
                "INCORRECT": "✗", "UNCERTAIN": "?"}.get(result["verdict"], " ")
        print(f"  [{icon}] {result['id']:6s}  {result['verdict']:<20s}  {q['note'][:55]}")
        results.append(result)

    tasks = [_run(q) for q in questions]
    await asyncio.gather(*tasks)
    return results


# ── Excel writer ──────────────────────────────────────────────────────────────

def _thin():
    s = Side(style="thin", color="FFAAAAAA")
    return Border(left=s, right=s, top=s, bottom=s)

def _fill(h: str) -> PatternFill:
    return PatternFill("solid", fgColor=h)

def _font(bold=False, color="FF000000", size=10) -> Font:
    return Font(bold=bold, color=color, size=size)

def _align(h="left", v="top", wrap=True) -> Alignment:
    return Alignment(horizontal=h, vertical=v, wrap_text=wrap)


def write_report(all_results: list[dict]) -> Path:
    wb = openpyxl.Workbook()

    # ── Sheet 1: All Results ──────────────────────────────────────────────────
    ws = wb.active
    ws.title = "Focused Regression"

    COLS = [
        ("Source",      18),
        ("ID",           8),
        ("Fix / Guard", 32),
        ("Question",    52),
        ("API Status",  16),
        ("Verdict",     18),
        ("Conf %",       8),
        ("Judge Reason",55),
        ("Rows",         7),
        ("Latency ms",  10),
        ("SQL",         70),
        ("Answer",      55),
    ]

    # Title row
    ws.merge_cells(f"A1:{chr(64+len(COLS))}1")
    ws["A1"] = "NEP Analytics — Focused Regression Report"
    ws["A1"].font      = Font(bold=True, color="FFFFFFFF", size=13)
    ws["A1"].fill      = _fill(C_DARK)
    ws["A1"].alignment = _align("center", "center")
    ws.row_dimensions[1].height = 22

    # Header row
    for ci, (h, w) in enumerate(COLS, 1):
        c = ws.cell(2, ci, h)
        c.font      = _font(bold=True, color="FFFFFFFF")
        c.fill      = _fill(C_MID)
        c.alignment = _align("center", "center", wrap=False)
        c.border    = _thin()
        ws.column_dimensions[c.column_letter].width = w
    ws.row_dimensions[2].height = 18
    ws.freeze_panes = "A3"

    # Group by source
    by_source: dict[str, list[dict]] = {}
    for r in all_results:
        by_source.setdefault(r["source"], []).append(r)

    row_idx = 3
    for src, rows in by_source.items():
        # Section separator
        ws.merge_cells(f"A{row_idx}:{chr(64+len(COLS))}{row_idx}")
        sep = ws.cell(row_idx, 1, f"  {src}")
        sep.font      = _font(bold=True, color="FFFFFFFF", size=10)
        sep.fill      = _fill(C_MID)
        sep.alignment = _align("left", "center", wrap=False)
        ws.row_dimensions[row_idx].height = 16
        row_idx += 1

        for res in rows:
            vf = VERDICT_FILL.get(res["verdict"], C_WHITE)
            af = C_ALT if row_idx % 2 == 0 else C_WHITE
            vals = [
                res["source"],
                res["id"],
                res["note"],
                res["question"],
                res["api_status"],
                res["verdict"],
                res["confidence"],
                res["reason"],
                res["row_count"] if res["row_count"] is not None else "—",
                res["latency_ms"],
                res["sql"],
                res["answer"],
            ]
            for ci, val in enumerate(vals, 1):
                cell = ws.cell(row_idx, ci, val)
                cell.border    = _thin()
                cell.alignment = _align(wrap=True)
                if ci == 6:  # Verdict
                    cell.fill = _fill(vf)
                    cell.font = _font(bold=True)
                    cell.alignment = _align("center", "center", wrap=False)
                else:
                    cell.fill = _fill(af)
                    cell.font = _font()
            ws.row_dimensions[row_idx].height = 85
            row_idx += 1

    # ── Sheet 2: Summary ──────────────────────────────────────────────────────
    ws2 = wb.create_sheet("Summary")

    ws2.merge_cells("A1:G1")
    ws2["A1"] = "Focused Regression — Summary"
    ws2["A1"].font      = Font(bold=True, color="FFFFFFFF", size=13)
    ws2["A1"].fill      = _fill(C_DARK)
    ws2["A1"].alignment = _align("center", "center")
    ws2.row_dimensions[1].height = 22

    hdrs = ["ID", "Fix / Guard", "Verdict", "Confidence", "API Status", "Rows", "Source"]
    for ci, h in enumerate(hdrs, 1):
        c = ws2.cell(3, ci, h)
        c.font      = _font(bold=True, color="FFFFFFFF")
        c.fill      = _fill(C_MID)
        c.border    = _thin()
        c.alignment = _align("center", "center", wrap=False)
    col_w = [8, 34, 20, 12, 18, 8, 20]
    for ci, w in enumerate(col_w, 1):
        ws2.column_dimensions[ws2.cell(3, ci).column_letter].width = w

    ri = 4
    for res in all_results:
        vf = VERDICT_FILL.get(res["verdict"], C_WHITE)
        vals = [
            res["id"],
            res["note"],
            res["verdict"],
            f"{res['confidence']}%",
            res["api_status"],
            res["row_count"] if res["row_count"] is not None else "—",
            res["source"],
        ]
        for ci, val in enumerate(vals, 1):
            c = ws2.cell(ri, ci, val)
            c.border    = _thin()
            c.alignment = _align("center", "center", wrap=False)
            if ci == 3:
                c.fill = _fill(vf)
                c.font = _font(bold=True)
            else:
                c.fill = _fill(C_ALT if ri % 2 == 0 else C_WHITE)
        ri += 1

    # Overall tallies
    ri += 1
    totals_label = ws2.cell(ri, 1, "OVERALL")
    totals_label.font = _font(bold=True, color="FFFFFFFF")
    totals_label.fill = _fill(C_MID)
    ws2.merge_cells(f"A{ri}:B{ri}")

    correct   = sum(1 for r in all_results if r["verdict"] == "CORRECT")
    partial   = sum(1 for r in all_results if r["verdict"] == "PARTIALLY_CORRECT")
    incorrect = sum(1 for r in all_results if r["verdict"] == "INCORRECT")
    uncertain = sum(1 for r in all_results if r["verdict"] == "UNCERTAIN")
    total     = len(all_results)

    summary_text = (
        f"CORRECT: {correct}/{total}   "
        f"PARTIAL: {partial}   "
        f"INCORRECT: {incorrect}   "
        f"UNCERTAIN: {uncertain}"
    )
    c = ws2.cell(ri, 3, summary_text)
    ws2.merge_cells(f"C{ri}:G{ri}")
    c.font = _font(bold=True)
    c.fill = _fill(C_ALT)
    c.alignment = _align("left", "center", wrap=False)

    wb.save(REPORT_PATH)
    return REPORT_PATH


# ── Main ──────────────────────────────────────────────────────────────────────

async def main():
    print("\nNEP Analytics — Focused Regression Runner")
    print(f"  Judge model : {JUDGE_MODEL}")
    print(f"  API         : {API_BASE}\n")

    # Load test data
    single_qs = load_single_turn_questions()
    mt_scens  = load_mt_scenarios()
    suite_qs  = load_suite_questions()

    print(f"  Plan: {len(single_qs)} single-turn  |  "
          f"{len(mt_scens)} MT scenarios  |  "
          f"{len(suite_qs)} suite questions\n")

    all_results: list[dict] = []

    # Single-turn
    print("── 78-question set ──────────────────────────────────────────────────")
    r1 = await run_single_turn(single_qs)
    all_results.extend(r1)

    # Multi-turn
    print("\n── Multi-turn scenarios ─────────────────────────────────────────────")
    r2 = await run_multi_turn(mt_scens)
    all_results.extend(r2)

    # Suite
    print("\n── 24-question suite ────────────────────────────────────────────────")
    r3 = await run_suite(suite_qs)
    all_results.extend(r3)

    # Sort by source order
    source_order = {"78-question set": 0, "Multi-turn": 1, "24-question suite": 2}
    all_results.sort(key=lambda r: (source_order.get(r["source"], 9), r["id"]))

    # Console summary
    correct   = sum(1 for r in all_results if r["verdict"] == "CORRECT")
    partial   = sum(1 for r in all_results if r["verdict"] == "PARTIALLY_CORRECT")
    incorrect = sum(1 for r in all_results if r["verdict"] == "INCORRECT")
    uncertain = sum(1 for r in all_results if r["verdict"] == "UNCERTAIN")
    total     = len(all_results)

    print(f"\n{'='*60}")
    print(f"  Focused Regression — {total} tests")
    print(f"{'='*60}")
    print(f"  ✓ CORRECT          : {correct} / {total}")
    print(f"  ~ PARTIALLY_CORRECT: {partial}")
    print(f"  ✗ INCORRECT        : {incorrect}")
    print(f"  ? UNCERTAIN        : {uncertain}")

    print(f"\n  Fix targets:")
    fix_ids = {"B3", "C1", "D5", "E1", "Q39"} | {f"MT10-T{i}" for i in range(1,5)}
    for r in all_results:
        if any(fi in r["id"] for fi in ["B3","C1","D5","E1","Q39","MT10"]):
            icon = {"CORRECT":"✓","PARTIALLY_CORRECT":"~","INCORRECT":"✗","UNCERTAIN":"?"}.get(r["verdict"]," ")
            print(f"    [{icon}] {r['id']:8s}  {r['verdict']}")

    out = write_report(all_results)
    print(f"\n  Report → {out}")
    print(f"{'='*60}\n")


if __name__ == "__main__":
    asyncio.run(main())
