"""
NEP Analytics — LLM Judge on NEP_Suite_Live_Results.xlsx
Reads the 24 live results from the Excel file, judges each one,
and writes NEP_Suite_Judge_Report.xlsx.

Usage (from repo root):
    python nep_analytics/tests/run_suite_judge_only.py
"""
from __future__ import annotations

import asyncio
import json
import sys
from pathlib import Path

import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

HERE      = Path(__file__).resolve().parent
REPO_ROOT = HERE.parent.parent
sys.path.insert(0, str(REPO_ROOT))

from nep_analytics.core.config import get_anthropic  # noqa: E402

LIVE_XLSX   = HERE / "NEP_Suite_Live_Results.xlsx"
JUDGE_XLSX  = HERE / "NEP_Suite_Judge_Report.xlsx"
JUDGE_MODEL = "claude-opus-5"
MAX_CONCUR  = 4

# ── Palette ───────────────────────────────────────────────────────────────────
C_DARK  = "002C5282"
C_MID   = "001E3A5F"
C_LIGHT = "00EBF5FF"
C_ALT   = "00F7F9FC"
C_WHITE = "00FFFFFF"
C_PASS  = "00D4EDDA"
C_WARN  = "00FFF3CD"
C_FAIL  = "00F8D7DA"
C_INFO  = "00D1ECF1"
C_GREY  = "00E9ECEF"

VERDICT_FILL = {
    "CORRECT":           C_PASS,
    "PARTIALLY_CORRECT": C_WARN,
    "INCORRECT":         C_FAIL,
    "UNCERTAIN":         C_INFO,
}

BEHAVIOR_LABEL = {
    "rows":            "Happy Path",
    "grain":           "Grain Trap",
    "normalize":       "Normalization",
    "correctly_empty": "Correctly Empty",
    "unanswerable":    "Unanswerable",
}

# ── Read live results from Excel ──────────────────────────────────────────────

def load_live_results() -> list[dict]:
    wb = openpyxl.load_workbook(LIVE_XLSX)
    ws = wb["Live Results"]
    results = []
    current_behavior = ""

    for r in range(1, ws.max_row + 1):
        v1 = ws.cell(r, 1).value
        # Category separator row (merged, starts with spaces)
        if v1 and str(v1).startswith("  "):
            current_behavior = str(v1).strip()
            continue
        # Skip title, header, and empty rows
        if not v1 or str(v1).strip() in ("ID", "NEP Analytics — Live DB Fact Extraction"):
            continue

        cid      = str(ws.cell(r, 1).value or "").strip()
        category = str(ws.cell(r, 2).value or "").strip()
        question = str(ws.cell(r, 3).value or "").strip()
        status   = str(ws.cell(r, 4).value or "").strip()
        latency  = ws.cell(r, 5).value
        rtype    = str(ws.cell(r, 6).value or "").strip()
        row_cnt  = ws.cell(r, 7).value
        answer   = str(ws.cell(r, 8).value or "").strip()
        sql      = str(ws.cell(r, 9).value or "").strip()

        if not cid or len(cid) > 5:
            continue

        results.append({
            "id":            cid,
            "behavior_label": current_behavior,
            "category":      category,
            "question":      question,
            "status":        status,
            "latency_ms":    latency,
            "response_type": rtype,
            "row_count":     row_cnt,
            "answer":        answer,
            "sql_used":      sql,
        })
    return results


# ── LLM Judge ────────────────────────────────────────────────────────────────

JUDGE_SYSTEM = """\
You are an expert SQL and analytics evaluator for the NEP (National Entrepreneurship Program) platform.

You will be given:
- A natural-language analytics question
- The behavior category (what kind of test this is)
- The SQL the system generated
- The answer the system returned

Evaluate whether the system handled the question CORRECTLY and return ONLY valid JSON:
{
  "verdict": "CORRECT" | "PARTIALLY_CORRECT" | "INCORRECT" | "UNCERTAIN",
  "confidence_pct": <integer 0-100>,
  "reasoning": "<2-4 sentence explanation focusing on SQL logic and answer quality>"
}

Verdict rules:
  CORRECT           — SQL is logically sound for the question; answer accurately reflects the data.
  PARTIALLY_CORRECT — SQL is mostly right but has a minor flaw, or answer is slightly off.
  INCORRECT         — SQL has a structural error (wrong table, wrong filter, wrong grain, wrong join),
                      OR the answer misrepresents/misses the key result.
  UNCERTAIN         — Cannot determine correctness without running the query against the live DB.

Behavior-specific guidance:
  Happy Path        — Check: correct table, correct filters, correct aggregation level.
  Grain Trap        — Check: COUNT(DISTINCT userid/user_id) must be used, NOT COUNT(*).
  Normalization     — Check: correct enum values used (camelCase for sessiontype, lowercase for program,
                      exact medium spellings like 'Emailer' vs 'Email', correct gapkey values).
  Correctly Empty   — Check: if result is 0/empty, did the answer explain why clearly (not just blank)?
  Unanswerable      — Check: did the system either (a) correctly identify the limitation and explain it,
                      or (b) construct a reasonable proxy query and clearly disclose it as a proxy?
                      Generating SQL and returning data WITHOUT flagging the limitation = INCORRECT.

NEP schema quick reference:
- nep_mentor_profiles_sample_data: one row per mentor×program; always COUNT(DISTINCT user_id)
  program column values: 'liftoff' | 'liftoff-spark' | 'liftoff-propel' | 'ignite' | 'activate' | ...
  mentor_type: 'EXPERT' | 'MENTOR' | 'SERVICE_PROVIDER'
  user_status: 'ACTIVE' | 'PENDING'
  visibility: 'PUBLIC' | 'INTERNAL'
  deleted: BOOLEAN
- nep_liftoffx_data_sample: one row per user activity; COUNT(DISTINCT userid) for user counts
  activity_type: 'signup' | 'message' | 'session' | 'mentor' | 'resource' | 'visitors' | 'repeat visitors' |
                 'jounrney_explore' | 'introductory_video_reg_users'
- nep_master_live_events_data: one row per event×participant
  sessiontype: 'expertSession' | 'roundTable'  (camelCase, exact)
  participant_status: 'REGISTERED' | 'ATTENDED' | 'NOSHOW'
  gapkey: 'PitchMastery' | 'CompetitiveStrategy' | 'GrowthHacking' | etc. (PascalCase, no spaces)
  program_key: 'liftoff' | 'liftoff-spark' | 'liftoff-propel'
"""


async def judge_one(sem: asyncio.Semaphore, anthropic, result: dict) -> dict:
    behavior_label = result.get("behavior_label", result.get("category", ""))
    user_msg = f"""Question: {result['question']}
Behavior category: {behavior_label}

SQL generated:
{result['sql_used'][:1200] or 'none'}

Answer returned:
{result['answer'][:1200] or 'none (zero results)'}

HTTP status: 200
Row count: {result['row_count'] or 'n/a (scalar)'}
"""
    try:
        async with sem:
            resp = await asyncio.get_event_loop().run_in_executor(
                None,
                lambda: anthropic.messages.create(
                    model=JUDGE_MODEL,
                    max_tokens=1500,
                    system=JUDGE_SYSTEM,
                    messages=[{"role": "user", "content": user_msg}],
                ),
            )
        raw = next((b.text for b in resp.content if b.type == "text"), "").strip()
        # Strip markdown fences if present
        if raw.startswith("```"):
            raw = raw.split("```")[1]
            if raw.startswith("json"):
                raw = raw[4:]
        verdict_data = json.loads(raw.strip())
    except Exception as exc:
        verdict_data = {
            "verdict":        "UNCERTAIN",
            "confidence_pct": 0,
            "reasoning":      f"Judge error: {exc}",
        }
    return {**result, "judge": verdict_data}


# ── Excel writer ──────────────────────────────────────────────────────────────

def _thin():
    s = Side(style="thin", color="FFAAAAAA")
    return Border(left=s, right=s, top=s, bottom=s)

def _fill(h):   return PatternFill("solid", fgColor=h)
def _font(bold=False, color="FF000000", size=10): return Font(bold=bold, color=color, size=size)
def _align(h="left", v="top", wrap=True): return Alignment(horizontal=h, vertical=v, wrap_text=wrap)


def write_excel(results: list[dict]) -> Path:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Judge Results"

    cols = [
        ("ID",              6),
        ("Category",        18),
        ("Question",        46),
        ("Status",          13),
        ("Verdict",         14),
        ("Confidence",      11),
        ("SQL (excerpt)",   60),
        ("Answer (excerpt)", 55),
        ("Reasoning",       70),
    ]

    # Title
    ws.merge_cells(f"A1:{chr(64+len(cols))}1")
    ws["A1"] = "NEP Analytics — Suite LLM Judge Report"
    ws["A1"].font      = Font(bold=True, color="FFFFFFFF", size=13)
    ws["A1"].fill      = _fill(C_DARK)
    ws["A1"].alignment = _align("center", "center")
    ws.row_dimensions[1].height = 22

    # Column headers
    for ci, (h, w) in enumerate(cols, 1):
        cell = ws.cell(2, ci, h)
        cell.font      = _font(bold=True, color="FFFFFFFF")
        cell.fill      = _fill(C_MID)
        cell.alignment = _align("center", "center", wrap=False)
        cell.border    = _thin()
        ws.column_dimensions[cell.column_letter].width = w
    ws.row_dimensions[2].height = 18
    ws.freeze_panes = "A3"

    # Group by behavior
    by_behavior: dict[str, list[dict]] = {}
    for r in results:
        lbl = r.get("behavior_label") or r.get("category", "Other")
        by_behavior.setdefault(lbl, []).append(r)

    row_idx = 3
    for blabel, rows in by_behavior.items():
        # Category separator
        ws.merge_cells(f"A{row_idx}:{chr(64+len(cols))}{row_idx}")
        cell = ws.cell(row_idx, 1, f"  {blabel}")
        cell.font      = _font(bold=True, color="FFFFFFFF", size=10)
        cell.fill      = _fill(C_MID)
        cell.alignment = _align("left", "center", wrap=False)
        ws.row_dimensions[row_idx].height = 16
        row_idx += 1

        for res in rows:
            j           = res.get("judge", {})
            verdict     = j.get("verdict", "UNCERTAIN")
            verdict_fill = VERDICT_FILL.get(verdict, C_GREY)
            alt_fill     = C_ALT if row_idx % 2 == 0 else C_WHITE

            vals = [
                res["id"],
                res["category"],
                res["question"],
                res["status"],
                verdict,
                f"{j.get('confidence_pct', '?')}%",
                (res["sql_used"] or "")[:400],
                (res["answer"] or "")[:400],
                j.get("reasoning", ""),
            ]
            for ci, val in enumerate(vals, 1):
                cell = ws.cell(row_idx, ci, val)
                cell.border    = _thin()
                cell.alignment = _align(wrap=True)
                if ci == 5:   # Verdict
                    cell.fill = _fill(verdict_fill)
                    cell.font = _font(bold=True)
                    cell.alignment = _align("center", "center", wrap=False)
                elif ci == 4: # Status
                    s_fill = {"HAS_RESULTS": C_PASS, "ZERO_RESULTS": C_WARN,
                              "ERROR": C_FAIL}.get(res["status"], C_WHITE)
                    cell.fill = _fill(s_fill)
                    cell.font = _font(bold=False)
                    cell.alignment = _align("center", "center", wrap=False)
                else:
                    cell.fill = _fill(alt_fill)
                    cell.font = _font()
            ws.row_dimensions[row_idx].height = 75
            row_idx += 1

    # ── Summary sheet ─────────────────────────────────────────────────────────
    ws2 = wb.create_sheet("Summary")
    ws2.merge_cells("A1:F1")
    ws2["A1"] = "Suite Judge Report — Summary"
    ws2["A1"].font      = Font(bold=True, color="FFFFFFFF", size=13)
    ws2["A1"].fill      = _fill(C_DARK)
    ws2["A1"].alignment = _align("center", "center")
    ws2.row_dimensions[1].height = 22

    hdrs = ["Category", "Total", "CORRECT", "PARTIALLY_CORRECT", "INCORRECT", "UNCERTAIN"]
    for ci, h in enumerate(hdrs, 1):
        cell = ws2.cell(3, ci, h)
        cell.font      = _font(bold=True, color="FFFFFFFF")
        cell.fill      = _fill(C_MID)
        cell.border    = _thin()
        cell.alignment = _align("center", "center", wrap=False)
        ws2.column_dimensions[cell.column_letter].width = 22

    overall = {"CORRECT": 0, "PARTIALLY_CORRECT": 0, "INCORRECT": 0, "UNCERTAIN": 0}
    ri = 4
    for blabel, rows in by_behavior.items():
        counts = {"CORRECT": 0, "PARTIALLY_CORRECT": 0, "INCORRECT": 0, "UNCERTAIN": 0}
        for r in rows:
            v = r.get("judge", {}).get("verdict", "UNCERTAIN")
            counts[v] = counts.get(v, 0) + 1
            overall[v] = overall.get(v, 0) + 1
        row_vals = [blabel, len(rows),
                    counts["CORRECT"], counts["PARTIALLY_CORRECT"],
                    counts["INCORRECT"], counts["UNCERTAIN"]]
        for ci, val in enumerate(row_vals, 1):
            cell = ws2.cell(ri, ci, val)
            cell.border    = _thin()
            cell.alignment = _align("center", "center", wrap=False)
            cell.fill      = _fill(C_ALT if ri % 2 == 0 else C_WHITE)
            if ci == 3 and val: cell.fill = _fill(C_PASS)
            if ci == 4 and val: cell.fill = _fill(C_WARN)
            if ci == 5 and val: cell.fill = _fill(C_FAIL)
        ri += 1

    # Totals row
    total = len(results)
    for ci, val in enumerate(["TOTAL", total,
                               overall["CORRECT"], overall["PARTIALLY_CORRECT"],
                               overall["INCORRECT"], overall["UNCERTAIN"]], 1):
        cell = ws2.cell(ri, ci, val)
        cell.font   = _font(bold=True, color="FFFFFFFF")
        cell.fill   = _fill(C_DARK)
        cell.border = _thin()
        cell.alignment = _align("center", "center", wrap=False)

    wb.save(JUDGE_XLSX)
    return JUDGE_XLSX


# ── Main ──────────────────────────────────────────────────────────────────────

async def main():
    anthropic = get_anthropic()
    results   = load_live_results()
    print(f"\nNEP Suite — LLM Judge")
    print(f"  Loaded {len(results)} results from {LIVE_XLSX.name}")
    print(f"  Judge model: {JUDGE_MODEL}\n")
    print("── Judging ──────────────────────────────────────────────────────────")

    sem     = asyncio.Semaphore(MAX_CONCUR)
    tasks   = [judge_one(sem, anthropic, r) for r in results]
    judged  = []
    for coro in asyncio.as_completed(tasks):
        res  = await coro
        j    = res.get("judge", {})
        icon = {"CORRECT": "✓", "PARTIALLY_CORRECT": "~",
                "INCORRECT": "✗", "UNCERTAIN": "?"}.get(j.get("verdict", ""), "?")
        print(f"  [{icon}] {res['id']:4s}  {j.get('verdict','?'):<20}  "
              f"conf={j.get('confidence_pct','?'):>3}%  {res['question'][:55]}")
        judged.append(res)

    judged.sort(key=lambda r: r["id"])
    out = write_excel(judged)

    # Summary
    overall = {"CORRECT": 0, "PARTIALLY_CORRECT": 0, "INCORRECT": 0, "UNCERTAIN": 0}
    for r in judged:
        v = r.get("judge", {}).get("verdict", "UNCERTAIN")
        overall[v] = overall.get(v, 0) + 1

    print(f"\n{'='*60}")
    print(f"  Suite Judge — {len(judged)} questions")
    print(f"{'='*60}")
    total = len(judged)
    for v, cnt in overall.items():
        icon = {"CORRECT": "✓", "PARTIALLY_CORRECT": "~",
                "INCORRECT": "✗", "UNCERTAIN": "?"}.get(v, "?")
        if cnt:
            print(f"  [{icon}] {v:<22}: {cnt:3d}  ({cnt/total*100:.0f}%)")

    print(f"\n  By category:")
    by_beh: dict[str, dict] = {}
    for r in judged:
        lbl = r.get("behavior_label") or r.get("category", "Other")
        bc  = by_beh.setdefault(lbl, {"CORRECT": 0, "PARTIALLY_CORRECT": 0,
                                       "INCORRECT": 0, "UNCERTAIN": 0, "total": 0})
        v = r.get("judge", {}).get("verdict", "UNCERTAIN")
        bc[v] += 1
        bc["total"] += 1
    for lbl, bc in by_beh.items():
        print(f"    {lbl:<22}  correct={bc['CORRECT']}  partial={bc['PARTIALLY_CORRECT']}"
              f"  incorrect={bc['INCORRECT']}  total={bc['total']}")
    print(f"\n  Report → {out}")
    print(f"{'='*60}\n")


if __name__ == "__main__":
    asyncio.run(main())
